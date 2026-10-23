"""GOAL4 §4 (decision 008): two-way sync with a person's EXISTING Google Sheet, against the fake Google.

The sheet already exists (never created by the app); the app writes only the person's own columns plus one hidden
key column; Sheet edits reach the app, app edits reach the right cell; last writer wins and is logged; rows are
never deleted on either side; formatting and dropdowns stay."""
import json
import sqlite3
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest import mock

import openpyxl

from job_hunt_agent import db, gsync, gsync_cycle, gsync_link, trackapi, tracker, tracker_map
from job_hunt_agent.google import auth, runner, setup, state
from job_hunt_agent.gsync_values import canon, serial, to_cell
from tests.fake_sheets import wrapped_writes
from tests.google_case import GoogleCase

SID = "1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789"
URL = f"https://docs.google.com/spreadsheets/d/{SID}/edit#gid=0"
TAB = "Job_Application_Tracker"
DROPDOWN = {9: {"condition": {"type": "ONE_OF_LIST"}}}               # "Email Sent"
ROWS = [["Acme", "ML Engineer", "09/12/2025", "JD one", "acme.pdf", "Ann Lee", "Recruiter", "Sent", "ann@acme.com",
         "Yes", 12, None, None, None],
        ["Beta", "Data Scientist", {"v": serial(datetime(2026, 5, 12)), "fmt": "DATE"}, "JD two", None, None, None,
         None, None, None, None, None, None, "Rejected"],
        ["Gamma", "AI Engineer", "10/01/2026", "JD three", None, None, None, None, None, None, None, None, None, None]]


def as_xlsx(rows: list[list], path: Path, header=tracker.COLUMNS, title=TAB) -> Path:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(list(header))
    for row in rows:
        sheet.append([(datetime(1899, 12, 30) + timedelta(days=v["v"])) if isinstance(v, dict) else v for v in row])
    book.save(path)
    return path


class SyncCase(GoogleCase):
    def setUp(self):
        super().setUp()
        self.book = self.fake.existing(SID, {TAB: [list(tracker.COLUMNS), *ROWS]})
        self.sheet = self.book.sheet(TAB)
        self.sheet["validation"].update(DROPDOWN)
        self.book.cell(TAB, 1, 0)["bold"] = True                                    # formatting that must stay
        tracker.import_tracker(as_xlsx(ROWS, self.root / "real.xlsx"), "jane", conn=self.conn)
        self.ids = {r["company"]: r["id"] for r in self.conn.execute("SELECT id, company FROM applications")}

    def tearDown(self):
        # user rule (2026-10-06 evening): no sync write ever puts wrapped or multi-line text into a visible cell
        try:
            self.assertEqual(wrapped_writes(self.book), [], "a write would wrap or span lines in a visible cell")
        finally:
            super().tearDown()

    def link(self):
        return gsync_link.link(self.conn, "jane", URL, apply=True)

    def grid(self) -> list[list]:
        return self.book.raw_values(TAB)

    def person_cells(self) -> list[list]:
        """The person's 14 columns of every row, empty cells as ''."""
        return [(row[:14] + [""] * 14)[:14] for row in self.grid()]

    def column(self, header: str) -> int:
        return self.grid()[0].index(header)

    def row_of(self, company: str) -> int:
        return next(i for i, r in enumerate(self.grid()) if r and r[0] == company)

    def edit(self, company: str, header: str, value) -> None:
        with self.fake.lock:
            self.book.cell(TAB, self.row_of(company), self.column(header)).update(
                value if isinstance(value, dict) else {"v": value, "formula": None})

    def sync(self) -> dict:
        return gsync_cycle.sync_profile(self.conn, "jane")[TAB]

    def app(self, company: str) -> dict:
        return dict(self.conn.execute("SELECT * FROM applications WHERE id = ?", (self.ids[company],)).fetchone())

    def log(self, kind: str) -> list[dict]:
        return [dict(r) for r in self.conn.execute("SELECT * FROM gsync_log WHERE kind = ?", (kind,))]


class LinkTest(SyncCase):
    def test_dry_run_writes_nothing_and_says_what_it_would_do(self):
        plan = gsync_link.link(self.conn, "jane", URL)
        self.assertFalse(plan["applied"])
        text = "\n".join(plan["lines"])
        self.assertIn("rows matched to their application: 3", text)
        self.assertIn('hidden key column: O "job-hunt-agent id" (added and hidden)', text)
        self.assertEqual(self.fake.count("POST", r":batchUpdate"), 0)
        self.assertIsNone(state.get(self.conn, "gsync.jane.sid"))

    def test_link_adds_one_hidden_key_column_and_changes_nothing_else(self):
        before = self.person_cells()
        self.link()
        grid = self.grid()
        self.assertEqual(grid[0][14], gsync.KEY_HEADER)
        self.assertEqual(self.sheet["hidden"], {14})
        self.assertEqual([r[14] for r in grid[1:4]], [f"jh:{self.ids[c]}" for c in ("Acme", "Beta", "Gamma")])
        self.assertEqual(self.person_cells(), before)                             # the person's cells untouched
        self.assertEqual(self.sheet["validation"], DROPDOWN)
        self.assertTrue(self.book.cell(TAB, 1, 0)["bold"])
        self.assertEqual(self.conn.execute("SELECT count(*) FROM gsync_rows").fetchone()[0], 3)

    def test_no_new_spreadsheet_is_ever_created(self):
        self.link()
        self.sync()
        self.assertEqual(list(self.fake.sheets), [SID])
        self.assertEqual(self.fake.count("POST", r"/drive/v3/files"), 0)

    def test_link_matches_a_moved_row_by_company_and_role(self):
        with self.fake.lock:
            grid = self.sheet["grid"]
            grid[1], grid[3] = grid[3], grid[1]
        self.link()
        grid = self.grid()
        self.assertEqual(grid[1][14], f"jh:{self.ids['Gamma']}")
        self.assertEqual(grid[3][14], f"jh:{self.ids['Acme']}")

    def test_cells_that_differ_at_link_are_listed_and_left_alone(self):
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        plan = gsync_link.link(self.conn, "jane", URL)
        self.assertIn("cells that differ between the Sheet and the app (left as they are): 1",
                      "\n".join(plan["lines"]))
        self.link()
        self.assertEqual(self.grid()[1][6], "Recruiter")                            # not pushed at link
        self.assertEqual(self.app("Acme")["target_role"], "Head of Talent")        # not pulled either

    def test_status_reports_cells_left_differing_since_linking(self):
        from job_hunt_agent.google import cli as gcli
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        self.link()
        self.assertEqual(runner.health(self.conn)["diverged"], {"jane": {"Target Role": 1}})
        self.assertIn("cells that differ between Sheet and app (left as they are): jane Target Role 1",
                      gcli.sync_line(self.conn))

    def test_the_page_header_never_reads_the_cells(self):
        # GOAL30-app: the page never shows the cell diff, and reading every row's cells (17 MB live) cost each
        # /api/google ~350 ms cold; `job-hunt-agent google status` still reports it
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        self.link()
        with mock.patch.object(runner, "diverged", side_effect=AssertionError("the header read the cells")):
            found = runner.banner(self.conn)
        self.assertNotIn("diverged", found["sync"])
        self.assertEqual(found["sync"]["linked"], ["jane"])
        self.assertEqual(runner.health(self.conn)["diverged"], {"jane": {"Target Role": 1}})
        plan = " ".join(r[3] for r in self.conn.execute(
            "EXPLAIN QUERY PLAN SELECT count(*) FROM gsync_rows WHERE missing_since IS NOT NULL"))
        self.assertIn("COVERING INDEX gsync_missing", plan)          # the count never reads the cells either

    def test_an_approved_fix_is_pushed_only_on_request(self):
        tracker.edit_field(self.conn, self.ids["Gamma"], "date_applied", "10/02/2026")
        self.link()
        self.sync()
        self.assertEqual(self.grid()[3][2], "10/01/2026")                           # left alone by the sync
        plan = gsync_link.push_diverged(self.conn, "jane", "Date Applied")
        self.assertEqual(plan["count"], 1)
        self.assertEqual(self.grid()[3][2], "10/01/2026")                           # preview changes nothing
        gsync_link.push_diverged(self.conn, "jane", "Date Applied", apply=True)
        self.assertEqual(self.grid()[3][2], "10/02/2026")
        self.assertEqual(self.sync().get("pushed", 0), 0)

    def test_a_date_cell_stays_a_date_when_the_app_fix_is_pushed(self):
        tracker.edit_field(self.conn, self.ids["Beta"], "date_applied", "05/14/2026")    # app holds text
        self.link()
        gsync_link.push_diverged(self.conn, "jane", "Date Applied", apply=True)
        cell = self.book.cell(TAB, 2, 2)
        self.assertEqual((cell["v"], cell.get("fmt")), (serial(datetime(2026, 5, 14)), "DATE"))

    def test_sheet_wins_pull_diverged_previews_then_pulls_every_column_and_never_writes_the_sheet(self):
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        tracker.set_status(self.conn, self.ids["Gamma"], "rejected")                # app-only status, sheet blank
        self.link()
        before, writes = self.person_cells(), self.fake.count("POST", r":batchUpdate")
        plan = gsync_link.pull_diverged(self.conn, "jane")
        self.assertEqual(plan["count"], 2)
        self.assertEqual(self.app("Acme")["target_role"], "Head of Talent")        # preview changes nothing
        gsync_link.pull_diverged(self.conn, "jane", apply=True)
        self.assertEqual(self.app("Acme")["target_role"], "Recruiter")
        self.assertEqual(self.app("Gamma")["status"], "applied")                   # blank Status: default pipeline
        self.assertEqual(runner.health(self.conn)["diverged"], {})
        self.assertEqual(self.person_cells(), before)
        self.assertEqual(self.fake.count("POST", r":batchUpdate"), writes)          # the sheet is never written
        self.assertEqual(self.sync().get("pushed", 0), 0)

    def test_compare_reads_every_cell_and_writes_nothing(self):
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        self.link()
        writes = self.fake.count("POST", r":batchUpdate")
        found = gsync_link.compare(self.conn, "jane")
        self.assertEqual(found["tabs"][TAB]["keyed"], 3)
        self.assertEqual([d[1:3] for d in found["differences"]], [(2, "Target Role")])
        gsync_link.pull_diverged(self.conn, "jane", apply=True)
        self.assertEqual(gsync_link.compare(self.conn, "jane")["differences"], [])
        self.assertEqual(self.fake.count("POST", r":batchUpdate"), writes)

    def test_pull_diverged_one_column_leaves_the_others(self):
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", "Head of Talent")
        tracker.edit_field(self.conn, self.ids["Acme"], "target_name", "Bo Chen")
        self.link()
        gsync_link.pull_diverged(self.conn, "jane", "Target Role", apply=True)
        self.assertEqual(self.app("Acme")["target_role"], "Recruiter")
        self.assertEqual(runner.health(self.conn)["diverged"], {"jane": {"Target Name": 1}})

    def test_spreadsheet_id_from_any_sheets_link(self):
        self.assertEqual(gsync.spreadsheet_id(URL), SID)
        self.assertEqual(gsync.spreadsheet_id(SID), SID)
        with self.assertRaises(ValueError):
            gsync.spreadsheet_id("https://example.com/not-a-sheet")


class TwoWayTest(SyncCase):
    def setUp(self):
        super().setUp()
        self.link()

    def test_sheet_edit_reaches_the_app(self):
        self.edit("Acme", "Target Name", "Ann Lee, Bo Chen")
        self.assertEqual(self.sync()["pulled"], 1)
        app = self.app("Acme")
        self.assertEqual(app["target_name"], "Ann Lee, Bo Chen")
        names = [c["name"] for c in self.conn.execute("SELECT name FROM contacts WHERE app_id = ? ORDER BY pos",
                                                      (app["id"],))]
        self.assertEqual(names, ["Ann Lee", "Bo Chen"])
        self.assertEqual(self.sync().get("pulled", 0), 0)                            # once, not every cycle

    def test_status_typed_in_the_sheet_moves_the_pipeline(self):
        self.edit("Gamma", "Status(A/R)", "Rejected")
        self.sync()
        self.assertEqual(self.app("Gamma")["status"], "rejected")

    def test_date_edited_in_the_sheet_moves_applied_on(self):
        self.edit("Beta", "Date Applied", {"v": serial(datetime(2026, 5, 14)), "fmt": "DATE"})
        self.sync()
        self.assertEqual(self.app("Beta")["applied_on"], "2026-05-14")

    def test_app_edit_writes_only_that_cell_and_keeps_its_format(self):
        cell = self.book.cell(TAB, self.row_of("Acme"), self.column("Email Address"))
        cell["fmt"] = "TEXT_FMT"
        trackapi.update(self.conn, "jane", self.ids["Acme"], {"email_address": "ann@acme.io"})
        before = [list(r) for r in self.grid()]
        self.assertEqual(self.sync()["pushed"], 1)
        after = self.grid()
        changed = [(r, c) for r, row in enumerate(after) for c, v in enumerate(row) if before[r][c] != v]
        self.assertEqual(changed, [(self.row_of("Acme"), self.column("Email Address"))])
        self.assertEqual(after[self.row_of("Acme")][self.column("Email Address")], "ann@acme.io")
        self.assertEqual(cell["fmt"], "TEXT_FMT")                                   # only userEnteredValue written
        self.assertEqual(self.sheet["validation"], DROPDOWN)

    def test_status_change_in_the_app_shows_in_status_column(self):
        tracker.set_status(self.conn, self.ids["Acme"], "interview")
        self.sync()
        self.assertEqual(self.grid()[self.row_of("Acme")][self.column("Status(A/R)")], "Interview")

    def test_a_sorted_sheet_still_gets_each_write_on_its_own_row(self):
        with self.fake.lock:
            grid = self.sheet["grid"]
            grid[1:4] = [grid[3], grid[1], grid[2]]                                 # the phone sorted the rows
        tracker.edit_field(self.conn, self.ids["Beta"], "target_name", "Dee")
        self.sync()
        self.assertEqual(self.grid()[self.row_of("Beta")][self.column("Target Name")], "Dee")
        self.assertEqual(self.grid()[self.row_of("Gamma")][self.column("Target Name")], "")

    def test_formula_looking_text_stays_text(self):
        tracker.edit_field(self.conn, self.ids["Acme"], "target_role", '=IMPORTXML("x")')
        self.sync()
        cell = self.book.cell(TAB, self.row_of("Acme"), self.column("Target Role"))
        self.assertIsNone(cell.get("formula"))
        self.assertEqual(cell["v"], '=IMPORTXML("x")')


class ConflictTest(SyncCase):
    def setUp(self):
        super().setUp()
        self.link()

    def both_edit(self, app_minutes_ago: int, read_minutes_ago: int = 10) -> None:
        now = datetime.now(timezone.utc)
        with db.tx(self.conn):
            state.put(self.conn, f"gsync.jane.{TAB}.read_at", db.iso(now - timedelta(minutes=read_minutes_ago)))
        tracker.edit_field(self.conn, self.ids["Acme"], "target_name", "App Value")
        stamps = json.dumps({"target_name": db.iso(now - timedelta(minutes=app_minutes_ago))})
        with db.tx(self.conn):
            self.conn.execute("UPDATE applications SET field_at = ? WHERE id = ?", (stamps, self.ids["Acme"]))
        self.edit("Acme", "Target Name", "Sheet Value")

    def test_sheet_later_wins_and_is_logged(self):
        self.both_edit(app_minutes_ago=8)                     # Sheet edit somewhere in the last 10 min: ~5 min ago
        self.assertEqual(self.sync()["conflicts"], 1)
        self.assertEqual(self.app("Acme")["target_name"], "Sheet Value")
        (row,) = self.log("conflict")
        self.assertEqual((row["winner"], row["sheet_value"], row["app_value"]), ("sheet", "Sheet Value", "App Value"))

    def test_app_later_wins_and_is_logged(self):
        self.both_edit(app_minutes_ago=1)
        self.sync()
        self.assertEqual(self.grid()[self.row_of("Acme")][self.column("Target Name")], "App Value")
        self.assertEqual(self.log("conflict")[0]["winner"], "app")


class RowsTest(SyncCase):
    def setUp(self):
        super().setUp()
        self.link()

    def test_row_typed_in_the_sheet_becomes_an_application(self):
        with self.fake.lock:
            for col, value in enumerate(["Delta", "ML Engineer", {"v": serial(datetime(2026, 10, 3)), "fmt": "DATE"}]):
                self.book.cell(TAB, 4, col).update(value if isinstance(value, dict) else {"v": value})
        self.assertEqual(self.sync()["new_from_sheet"], 1)
        app = dict(self.conn.execute("SELECT * FROM applications WHERE company = 'Delta'").fetchone())
        self.assertEqual((app["source"], app["applied_on"], app["role_title"]), ("sheet", "2026-10-03", "ML Engineer"))
        self.assertEqual(self.grid()[4][14], f"jh:{app['id']}")
        self.assertEqual(self.sync().get("new_from_sheet", 0), 0)                   # once

    def test_date_only_placeholder_rows_are_left_alone(self):
        with self.fake.lock:
            self.book.cell(TAB, 4, 2).update({"v": serial(datetime(2026, 2, 10)), "fmt": "DATE"})
        self.assertEqual(self.sync().get("new_from_sheet", 0), 0)
        self.assertEqual(self.grid()[4][14:], [])

    def test_applied_in_the_app_is_added_at_the_end_with_its_key(self):
        made = trackapi.create(self.conn, "jane", {"company": "Epsilon", "role_title": "AI Engineer",
                                                   "location": "Austin, TX", "status": "saved"})
        self.assertEqual(self.sync().get("appended", 0), 0)                         # a bare save stays off the sheet
        tracker.set_status(self.conn, made["id"], "applied")
        self.assertEqual(self.sync()["appended"], 1)
        row = self.grid()[4]
        self.assertEqual((row[0], row[1], row[14]), ("Epsilon", "AI Engineer", f"jh:{made['id']}"))
        self.assertEqual(self.sync().get("appended", 0), 0)

    def test_a_job_typed_in_both_places_is_paired_not_duplicated(self):
        made = trackapi.create(self.conn, "jane", {"company": "Aspenview Technology Partners",
                                                   "role_title": "Senior Agentic Engineer"})
        with self.fake.lock:
            for col, value in enumerate(["AspenView", "Agentic Engineer", datetime.now().strftime("%m/%d/%Y")]):
                self.book.cell(TAB, 4, col).update({"v": value})
        plan = gsync_link.plans(self.conn, "jane", SID)[0][0]
        self.assertEqual([(r.index, a["id"]) for r, a in plan["paired"]], [(4, made["id"])])
        report = self.sync()
        self.assertEqual((report["paired"], report.get("new_from_sheet", 0), report.get("appended", 0)), (1, 0, 0))
        self.assertEqual(self.grid()[4][14], f"jh:{made['id']}")
        self.assertEqual(self.grid()[4][1], "Agentic Engineer")                    # her wording kept
        self.assertEqual(self.conn.execute("SELECT count(*) FROM applications WHERE company LIKE 'Aspen%'")
                         .fetchone()[0], 1)

    def test_a_row_deleted_on_the_sheet_is_reported_never_deleted_or_re_added(self):
        with self.fake.lock:
            del self.sheet["grid"][2]                                              # Beta's row
        self.assertEqual(self.sync()["missing"], 1)
        self.assertIsNotNone(self.app("Beta"))
        self.sync()
        self.assertNotIn("Beta", [r[0] for r in self.grid() if r])
        self.assertEqual(runner.health(self.conn)["missing_rows"], 1)

    def test_a_copied_row_with_its_hidden_key_becomes_its_own_application(self):
        with self.fake.lock:
            self.sheet["grid"].append([dict(c) for c in self.sheet["grid"][1]])     # Acme copied, key and all
            self.book.cell(TAB, 4, 1)["v"] = "Staff ML Engineer"
        self.assertEqual(self.sync()["new_from_sheet"], 1)
        new_id = self.conn.execute("SELECT id FROM applications WHERE role_title = 'Staff ML Engineer'").fetchone()[0]
        self.assertEqual(self.grid()[4][14], f"jh:{new_id}")
        self.assertEqual(self.log("duplicate_key")[0]["sheet_value"], f"jh:{self.ids['Acme']}")


class AddColumnsTest(SyncCase):
    def setUp(self):
        super().setUp()
        self.add_profile("jane", "name: Jane Doe\nfamilies: [ml_engineer]\n"
                                 "tracker_add_columns: [{header: Address, after: Role Title}]\n")
        with db.tx(self.conn):
            self.conn.execute("UPDATE applications SET address = 'Austin, TX' WHERE id = ?", (self.ids["Gamma"],))
        self.link()

    def test_preview_changes_nothing(self):
        text = "\n".join(gsync_link.add_columns(self.conn, "jane")["lines"])
        self.assertIn('insert column C "Address" right after "Role Title"', text)
        self.assertIn("the next sync fills 1 Address cells", text)
        self.assertEqual(self.grid()[0][2], "Date Applied")

    def test_insert_moves_columns_dropdowns_and_key_then_fills_address(self):
        gsync_link.add_columns(self.conn, "jane", apply=True)
        grid = self.grid()
        self.assertEqual(grid[0][:4], ["Company Name", "Role Title", "Address", "Date Applied"])
        self.assertEqual(grid[0][15], gsync.KEY_HEADER)
        self.assertEqual(self.sheet["hidden"], {15})
        self.assertEqual(self.sheet["validation"], {10: DROPDOWN[9]})
        self.assertEqual(grid[self.row_of("Gamma")][2], "Austin, TX")
        self.assertEqual(grid[self.row_of("Acme")][2], "")
        self.assertEqual(grid[self.row_of("Acme")][3], "09/12/2025")              # its data moved with it


class MappedProfileTest(GoogleCase):
    """Jordan's shape: her own header spellings ("Addresss"), a pass-through "Column 1", her order."""

    HEADER = ["Company Name", "Role Title", "Date Applied", "Addresss", "Status", "Column 1"]
    TYPED_OCT_5 = {"v": serial(datetime(2026, 5, 10)), "fmt": "DATE"}   # "10/05/2026" typed into a dd/mm sheet

    def setUp(self):
        super().setUp()
        rows = [["Acme", "AI Engineer", self.TYPED_OCT_5, "Remote", None, "note"]]
        template = as_xlsx(rows, self.root / "jordan.xlsx", self.HEADER, "Sheet1")
        folder = self.add_profile("sri", "name: Sri\nfamilies: [ai_engineer]\n")
        (folder / "tracker_map.yaml").write_text(
            f"extras: [Address]\ntemplate: {template}\nsheets:\n  - name: Sheet1\n    columns:\n"
            "      Company Name: Company Name\n      Role Title: Role Title\n      Date Applied: Date Applied\n"
            "      Addresss: Address\n      Status: Status(A/R)\nswap_date_cells: true\n")
        self.conn.execute("INSERT INTO profiles(id, name) VALUES ('sri', 'Sri')")
        tracker_map.import_mapped(template, "sri", folder / "tracker_map.yaml", self.conn)
        self.book = self.fake.existing(SID, {"Sheet1": [self.HEADER, *rows]})
        gsync_link.link(self.conn, "sri", URL, apply=True)
        self.app_id = self.conn.execute("SELECT id FROM applications WHERE profile = 'sri'").fetchone()[0]

    def edit(self, col: int, value) -> None:
        with self.fake.lock:
            self.book.cell("Sheet1", 1, col).update({"v": value, "formula": None})

    def test_her_swapped_dates_are_the_same_date_both_ways(self):
        app = dict(self.conn.execute("SELECT * FROM applications WHERE id = ?", (self.app_id,)).fetchone())
        self.assertEqual(app["applied_on"], "2026-10-05")                          # the import's swap
        cells = json.loads(self.conn.execute("SELECT cells FROM gsync_rows WHERE app_id = ?",
                                             (self.app_id,)).fetchone()[0])
        self.assertEqual(cells["Date Applied"][0], cells["Date Applied"][1])        # not a difference
        tracker.edit_field(self.conn, self.app_id, "date_applied", "10/06/2026")
        gsync_cycle.sync_profile(self.conn, "sri")
        cell = self.book.cell("Sheet1", 1, 2)
        self.assertEqual(cell["v"], serial(datetime(2026, 6, 10)))                 # shows 10/06/2026 in her dd/mm

    def test_her_columns_her_order_and_the_key_after_them(self):
        self.assertEqual(self.book.raw_values("Sheet1")[0], [*self.HEADER, gsync.KEY_HEADER])

    def test_her_address_and_pass_through_edits_reach_the_app_and_back(self):
        self.edit(3, "Dallas, TX")
        self.edit(5, "called back")
        gsync_cycle.sync_profile(self.conn, "sri")
        app = dict(self.conn.execute("SELECT * FROM applications WHERE id = ?", (self.app_id,)).fetchone())
        self.assertEqual(app["address"], "Dallas, TX")
        raw = json.loads(app["raw_json"])
        self.assertEqual([c["v"] for c in raw if isinstance(c, dict) and c.get("col") == "Column 1"], ["called back"])
        trackapi.update(self.conn, "sri", self.app_id, {"address": "Plano, TX"})
        gsync_cycle.sync_profile(self.conn, "sri")
        self.assertEqual(self.book.raw_values("Sheet1")[1][3], "Plano, TX")
        self.assertEqual(self.book.raw_values("Sheet1")[1][5], "called back")


class StatusTest(SyncCase):
    def test_runner_syncs_linked_people_and_reports_health(self):
        self.assertIn("skipped", runner.run(self.conn))                             # nothing linked yet
        self.link()
        self.edit("Acme", "Target Role", "TA Lead")
        report = runner.run(self.conn)
        self.assertEqual(report["profiles"]["jane"][TAB]["pulled"], 1)
        health = runner.health(self.conn)
        self.assertTrue(health["healthy"])
        self.assertEqual(health["linked"], ["jane"])
        from job_hunt_agent.google import cli as gcli
        self.assertTrue(gcli.sync_line(self.conn).startswith("sheet sync: healthy - linked jane"))

    def test_a_failing_sheet_is_an_error_line_not_a_crash(self):
        self.link()
        self.fake.fail("GET", r"/v4/spreadsheets/", 403, times=10, reason="forbidden")
        report = runner.run(self.conn)
        self.assertIn("jane", report["errors"])
        self.assertFalse(runner.health(self.conn)["healthy"])
        self.assertIn("Sheet sync failed", runner.banner(self.conn)["message"])

    def test_a_bug_in_one_persons_cycle_is_an_error_line_not_a_crash(self):
        self.link()
        with mock.patch.object(runner.gsync_cycle, "sync_profile", side_effect=sqlite3.IntegrityError("boom")):
            report = runner.run(self.conn)
        self.assertIn("boom", report["errors"]["jane"])
        self.assertIn("jane", state.get_json(self.conn, "last_sync_errors"))

    def test_a_failed_cycle_keeps_the_edit_waiting_and_retries_it_soon(self):
        self.link()
        tracker.edit_field(self.conn, self.ids["Acme"], "target_name", "x")       # an app edit waiting to go out
        self.fake.fail("GET", r"/v4/spreadsheets/", 429, times=10, reason="quota")
        runner.run(self.conn)
        self.assertIsNotNone(state.dirty(self.conn)[0])                            # not dropped by the failure
        now = datetime.now(timezone.utc)
        self.assertFalse(runner.due(self.conn, now))                               # no hammering a quota limit
        self.assertTrue(runner.due(self.conn, now + timedelta(seconds=runner.RETRY_S)))
        self.assertLess(runner.RETRY_S + runner.TICK_S, runner.POLL_S)

    def test_due_after_an_app_edit_or_the_poll_interval(self):
        now = datetime.now(timezone.utc)
        with db.tx(self.conn):
            state.put(self.conn, "last_sync_at", db.iso(now))
        self.assertFalse(runner.due(self.conn, now))
        tracker.edit_field(self.conn, self.ids["Acme"], "target_name", "x")
        self.assertTrue(runner.due(self.conn, now + timedelta(seconds=runner.DEBOUNCE_S)))
        self.assertLessEqual(runner.DEBOUNCE_S + runner.TICK_S, 60)                 # app -> Sheet within a minute
        self.assertLessEqual(runner.POLL_S + runner.TICK_S, 300)                    # Sheet -> app within 5 minutes

    def test_production_client_has_no_seven_day_clock(self):
        token = {**auth.load(), "mode": "production"}
        auth.save(token)
        status = auth.status()
        self.assertIsNone(status["days_left"])
        self.assertFalse(status["warn"])
        auth.save({**token, "mode": "testing",
                   "refresh_obtained_at": db.iso(db.utcnow() - timedelta(days=6, hours=1))})
        self.assertTrue(auth.status()["warn"])

    def test_scopes_are_sheets_and_app_created_drive_files_and_setup_says_in_production(self):
        from job_hunt_agent.google import SCOPES
        self.assertEqual(SCOPES, ("https://www.googleapis.com/auth/spreadsheets",
                                  "https://www.googleapis.com/auth/drive.file"))
        self.assertIn("In production", setup.steps_text())
        self.assertIn("drive.file", setup.steps_text())


class ValuesTest(GoogleCase):
    def test_a_date_compares_equal_as_text_or_as_a_date_cell(self):
        self.assertEqual(canon("05/12/2026", "Date Applied"), canon(serial(datetime(2026, 5, 12)), "Date Applied"))
        self.assertNotEqual(canon("05/12/2026", "Target Name"), canon(serial(datetime(2026, 5, 12)), "Target Name"))

    def test_canon_makes_display_formats_irrelevant(self):
        self.assertEqual(canon(datetime(2026, 5, 12)), canon(serial(datetime(2026, 5, 12))))
        self.assertEqual(canon(12.0), canon(12))
        self.assertEqual(canon(None), canon(""))

    def test_to_cell_types(self):
        self.assertEqual(to_cell("=SUM(A1)", as_text=True)["userEnteredValue"], {"stringValue": "=SUM(A1)"})
        self.assertEqual(to_cell("=SUM(A1)", as_text=False)["userEnteredValue"], {"formulaValue": "=SUM(A1)"})
        self.assertEqual(to_cell(None, False)["userEnteredValue"], {})
        self.assertEqual(to_cell(datetime(2026, 5, 12), False, new_row=True)["userEnteredFormat"]["numberFormat"]
                         ["type"], "DATE")
        self.assertEqual(to_cell(datetime(2026, 5, 12), False)["userEnteredFormat"], {"wrapStrategy": "CLIP"})
        self.assertIn("textFormatRuns", to_cell("JD", False, "https://x.co", new_row=True))
        one = to_cell("12 Main St\r\n  Austin, TX", True, header="Address")         # one clipped line (2026-10-06)
        self.assertEqual((one["userEnteredValue"], one["userEnteredFormat"]),
                         ({"stringValue": "12 Main St Austin, TX"}, {"wrapStrategy": "CLIP"}))
        self.assertEqual(to_cell("a\nb", True, header="JD Text")["userEnteredValue"], {"stringValue": "a\nb"})
        self.assertEqual(canon("12 Main St\nAustin, TX", "Address"), canon("12 Main St Austin, TX", "Address"))
        self.assertNotEqual(canon("a\nb", "JD Text"), canon("a b", "JD Text"))
