"""each person's tracker in their real sheet's columns (plus Address for Alex), the app's own fields kept in a
separate section, and an export -> import round trip that loses nothing (sheet cells and app fields)."""
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from job_hunt_agent import config, db, trackapi, tracker, tracker_layout
from tests.test_tracker import make_fixture

REAL_PROFILE = config.profile
ADDRESS = [{"header": "Address", "after": "Role Title"}]
WITH_ADDRESS = [*tracker.COLUMNS[:2], "Address", *tracker.COLUMNS[2:]]


def with_address(pid: str) -> dict:
    """Alex's GOAL4 profile shape: the user's real columns plus Address after Role Title (template given per test)."""
    base = {k: v for k, v in REAL_PROFILE(pid).items() if k != "tracker_template"}
    return {**base, "tracker_add_columns": ADDRESS}


def dress(path: Path) -> None:
    """The real sheet's furniture: dropdowns on A2:J5 and L2:N5, a CF rule anchored on D1, a table, column widths."""
    book = openpyxl.load_workbook(path)
    sheet = book.active
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv.add("A2:J5")
    dv.add("L2:N5")
    sheet.add_data_validation(dv)
    sheet.conditional_formatting.add("D1:E20", FormulaRule(formula=["LEN(TRIM(D1))>0"],
                                                            fill=PatternFill("solid", start_color="FFEEEE")))
    sheet.add_table(Table(displayName="Job_Application_Tracker", ref="A1:N5"))
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["C"].width = 14
    sheet.freeze_panes = "A2"
    book.save(path)


def grid(path: Path) -> list[list]:
    return [[c.value for c in row] for row in openpyxl.load_workbook(path).worksheets[0].iter_rows()]


class InsertColumnTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / "real.xlsx"
        make_fixture(self.path)
        dress(self.path)

    def tearDown(self):
        self.tmp.cleanup()

    def test_insert_keeps_dropdowns_formatting_table_and_widths(self):
        book = openpyxl.load_workbook(self.path)
        tracker_layout.insert_column(book.active, 3, "Address", 30)
        out = Path(self.tmp.name) / "out.xlsx"
        book.save(out)
        sheet = openpyxl.load_workbook(out).active
        self.assertEqual([c.value for c in sheet[1]], WITH_ADDRESS)
        self.assertEqual(sorted(str(d.sqref) for d in sheet.data_validations.dataValidation), ["A2:K5 M2:O5"])
        (cf,) = list(sheet.conditional_formatting)
        self.assertEqual(str(cf.sqref), "E1:F20")
        self.assertEqual(cf.rules[0].formula, ["LEN(TRIM(E1))>0"])              # the formula anchor moved too
        table = sheet.tables["Job_Application_Tracker"]
        self.assertEqual(table.ref, "A1:O5")
        self.assertEqual([c.name for c in table.tableColumns][:4], WITH_ADDRESS[:4])
        self.assertEqual(sheet.column_dimensions["A"].width, 22)
        self.assertEqual(sheet.column_dimensions["C"].width, 30)                 # the new column
        self.assertEqual(sheet.column_dimensions["D"].width, 14)                 # old C moved right
        self.assertEqual(sheet["D2"].value, "09/12/2025")                        # data moved with its column
        self.assertIsNone(sheet["C2"].value)

    def test_shift_range_moves_widens_or_keeps(self):
        self.assertEqual(tracker_layout._shift_range("D1:E9", 3), "E1:F9")      # right of it: moves
        self.assertEqual(tracker_layout._shift_range("A2:J9", 3), "A2:K9")      # across it: widens
        self.assertEqual(tracker_layout._shift_range("A1:B9", 3), "A1:B9")      # left of it: unchanged


class PlainLayoutTest(unittest.TestCase):
    """Alex: the user's real workbook's columns in order, Address after Role Title, app fields apart; round trip lossless."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.real = self.dir / "real.xlsx"
        make_fixture(self.real)
        dress(self.real)
        self.conn = db.connect(self.dir / "t.db")
        self.patch = mock.patch.object(config, "profile", with_address)
        self.patch.start()
        tracker.import_tracker(self.real, "alex", conn=self.conn)

    def tearDown(self):
        self.patch.stop()
        self.conn.close()
        self.tmp.cleanup()

    def app_id(self, company: str) -> int:
        return self.conn.execute("SELECT id FROM applications WHERE company = ? ORDER BY id", (company,)).fetchone()[0]

    def export(self, name: str, conn=None) -> Path:
        return tracker.export_tracker(self.dir / name, "alex", template=self.real, conn=conn or self.conn)

    def test_layout_is_the_real_header_plus_address(self):
        layout = tracker_layout.of_profile("alex", self.real)
        (tab,) = layout.tabs
        self.assertEqual(tab.headers, WITH_ADDRESS)
        self.assertEqual(tab.name, "Job_Application_Tracker")
        described = tracker_layout.describe(layout)
        self.assertEqual(described["extras_label"], tracker_layout.EXTRAS_SHEET)
        self.assertNotIn("Pipeline status", [c["header"] for c in described["tabs"][0]["columns"]])
        self.assertIn("Pipeline status", [x["label"] for x in described["extras"]])

    def test_view_shows_sheet_cells_in_order_and_extras_apart(self):
        app_id = self.app_id("Other AI")
        trackapi.update(self.conn, "alex", app_id, {"address": "Austin, TX", "notes": "ping in May"})
        app = dict(self.conn.execute("SELECT * FROM applications WHERE id = ?", (app_id,)).fetchone())
        view = trackapi.view(app, [], date.today(), tracker_layout.of_profile("alex", self.real))
        cells = view["sheet"]["cells"]
        self.assertEqual(len(cells), 15)
        self.assertEqual(cells[:4], ["Other AI", "AI Engineer – Agents", "Austin, TX", "12/15/2025"])
        self.assertEqual(view["sheet"]["tab"], "Job_Application_Tracker")
        self.assertEqual((view["extras"]["notes"], view["extras"]["status"]), ("ping in May", "applied"))
        self.assertNotIn("ping in May", cells)                                   # app fields never mix into the sheet

    def test_export_has_address_column_furniture_and_extras_sheet(self):
        book = openpyxl.load_workbook(self.export("out.xlsx"))
        self.assertEqual(book.sheetnames, ["Job_Application_Tracker", tracker_layout.EXTRAS_SHEET])
        sheet = book["Job_Application_Tracker"]
        self.assertEqual([c.value for c in sheet[1]], WITH_ADDRESS)
        self.assertEqual(sorted(str(d.sqref) for d in sheet.data_validations.dataValidation), ["A2:K5 M2:O5"])
        self.assertEqual(sheet.tables["Job_Application_Tracker"].ref, "A1:O5")
        self.assertEqual(sheet.freeze_panes, "A2")
        extras = book[tracker_layout.EXTRAS_SHEET]
        header = [c.value for c in extras[1]]
        self.assertEqual(header[:5], ["Sheet tab", "Sheet row", "Row ID", "Company Name", "Role Title"])
        self.assertIn("Pipeline status", header)
        self.assertEqual(extras.max_row - 1, 4)

    def test_roundtrip_into_an_empty_db_keeps_sheet_cells_and_app_fields(self):
        trackapi.update(self.conn, "alex", self.app_id("Other AI"),
                        {"address": "Austin, TX", "notes": "ping", "job_url": "https://x.co/1", "salary": "$150k"})
        first = self.export("one.xlsx")
        fresh = db.connect(self.dir / "fresh.db")
        try:
            res = tracker.import_tracker(first, "alex", conn=fresh)
            self.assertEqual(res["inserted"], 4)
            self.assertGreaterEqual(res["extras_restored"], 1)
            got = dict(fresh.execute("SELECT * FROM applications WHERE company = 'Other AI'").fetchone())
            self.assertEqual((got["address"], got["notes"], got["job_url"], got["salary"]),
                             ("Austin, TX", "ping", "https://x.co/1", "$150k"))
            second = self.export("two.xlsx", fresh)
        finally:
            fresh.close()
        self.assertEqual(grid(first), grid(second))

    def test_reimport_of_own_export_never_clears_an_app_field(self):
        app_id = self.app_id("Recent Inc")
        trackapi.update(self.conn, "alex", app_id, {"notes": "keep me"})
        out = self.export("one.xlsx")
        book = openpyxl.load_workbook(out)
        extras = book[tracker_layout.EXTRAS_SHEET]
        notes = [c.value for c in extras[1]].index("Notes") + 1
        for r in range(2, extras.max_row + 1):
            extras.cell(r, notes).value = None                                   # a blank cell is "no value"
        book.save(out)
        tracker.import_tracker(out, "alex", conn=self.conn)
        self.assertEqual(self.conn.execute("SELECT notes FROM applications WHERE id = ?", (app_id,)).fetchone()[0],
                         "keep me")


class CellTest(unittest.TestCase):
    def test_passthrough_column_comes_from_the_tagged_import_cell(self):
        column = tracker_layout.Column("Column 1", None)
        layout = tracker_layout.Layout("jordan", (tracker_layout.Tab("Sheet1", "apps", (column,)),), None, None)
        app = {"raw_json": json.dumps([{"t": "str", "v": "Acme"}, {"col": "Column 1", "t": "str", "v": "Rejected"}])}
        self.assertEqual(tracker_layout.cell(app, column, layout), ("Rejected", None, False))
        self.assertIsNone(tracker_layout.cell({"raw_json": None}, column, layout)[0])

    def test_formula_looking_text_is_written_as_text(self):
        self.assertTrue(tracker_layout._as_text("=SUM(A1)"))
        self.assertFalse(tracker_layout._as_text("plain"))
        self.assertFalse(tracker_layout._as_text(12))

    def test_read_extras_skips_rows_without_a_key_and_never_imports_status_or_timestamps(self):
        book = openpyxl.Workbook()
        tracker_layout.write_extras(book, [("Tab", 2, {"id": 9, "company": "A", "notes": "n", "status": "applied",
                                                       "updated_at": "2026-10-05T00:00:00Z", "source": "feed"})])
        book[tracker_layout.EXTRAS_SHEET].append([None, None, None, "orphan"])
        self.assertEqual(tracker_layout.read_extras(book), {("Tab", 2): {"notes": "n", "source": "feed"}})


if __name__ == "__main__":
    unittest.main()
