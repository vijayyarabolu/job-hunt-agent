"""facts a vendor model reads from a resume file: total years from its date ranges and the layout features
vendors document as parsing risks (tables, text boxes, header/footer text, columns, images, file type and size)."""
import io
import unittest
from datetime import date

import docx

from job_hunt_agent import resume_facts

TODAY = date(2026, 10, 6)


def _docx(build) -> bytes:
    document = docx.Document()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


class YearsTest(unittest.TestCase):
    def test_month_ranges_and_present_are_summed(self):
        text = "IBM, India  Jan 2021 – Jul 2023\nDeloitte  Aug 2023 - Present"
        self.assertAlmostEqual(resume_facts.years_of(text, TODAY), 70 / 12, places=2)        # 31 + 39 months, end month counted

    def test_overlapping_ranges_count_once(self):
        text = "A  Jan 2020 – Dec 2021\nB  Jun 2021 – Dec 2022"
        self.assertAlmostEqual(resume_facts.years_of(text, TODAY), 3.0, places=2)

    def test_year_only_ranges_and_current(self):
        self.assertAlmostEqual(resume_facts.years_of("Acme 2019 - 2021\nBeta 2024 – Current", TODAY), 58 / 12, places=2)

    def test_education_dates_are_not_experience(self):
        text = "Work Experience\nAcme  Jan 2022 – Jan 2024\nEducation\nMS Computer Science  Aug 2023 – May 2025"
        self.assertAlmostEqual(resume_facts.years_of(text, TODAY), 25 / 12, places=2)

    def test_no_dates_is_none(self):
        self.assertIsNone(resume_facts.years_of("Python, SQL, Spark", TODAY))


class RisksTest(unittest.TestCase):
    def test_a_plain_docx_has_no_risks(self):
        data = _docx(lambda d: d.add_paragraph("Experience"))
        self.assertEqual(resume_facts.parse_risks(data, "cv.docx"), ())

    def test_table_and_header_text_are_found(self):
        def build(d):
            d.add_table(rows=1, cols=2)
            d.sections[0].header.paragraphs[0].text = "Jane Doe · jane@example.com"
        risks = resume_facts.parse_risks(_docx(build), "cv.docx")
        self.assertEqual(set(risks), {"tables", "header_footer_text"})

    def test_unreadable_bytes_and_unknown_type(self):
        self.assertEqual(resume_facts.parse_risks(b"not a zip", "cv.docx"), ("unreadable",))
        self.assertIn("file_type", resume_facts.parse_risks(b"x", "cv.pages"))

    def test_size_is_reported_in_megabytes(self):
        self.assertEqual(resume_facts.size_mb(b"x" * 3 * 1024 * 1024), 3.0)


if __name__ == "__main__":
    unittest.main()
