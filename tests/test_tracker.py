"""Tracker import/export round-trip: synthetic fixture always; the real 374-row file when present locally."""
import copy
import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest import mock

import openpyxl

from job_hunt_agent import config, db, tracker

REAL = config.DATA / "import" / "Application_Tracker.xlsx"


def read_cells(path: Path) -> list[list]:
    """The 14 tracker columns by header name (an export may add columns, e.g. GOAL4's Address), in COLUMNS order."""
    sheet = openpyxl.load_workbook(path).worksheets[0]
    header = [c.value for c in sheet[1]]
    at = [header.index(c) for c in tracker.COLUMNS]
    rows = []
    for row in sheet.iter_rows(min_row=2, max_col=len(header), values_only=True):
        picked = [row[i] for i in at]
        if any(v not in (None, "") for v in picked):
            rows.append(picked)
    return rows


def same_cell(a, b) -> bool:
    """Excel has one numeric type: 9.0 written back reads as int 9. Everything else must match type exactly."""
    numeric = (int, float)
    if isinstance(a, numeric) and isinstance(b, numeric) and not isinstance(a, bool) and not isinstance(b, bool):
        return a == b
    return a == b and type(a) is type(b)


def make_fixture(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Job_Application_Tracker"
    sheet.append(tracker.COLUMNS)
    recent = (date.today() - timedelta(days=4)).strftime("%m/%d/%Y")
    sheet.append(["Example Co", "ML Engineer", "09/12/2025", "JD text", "v1", None, None, "Sent", "N/A", "No",
                  12.0, None, None, None])
    sheet.append(["Other AI", "AI Engineer – Agents", datetime(2025, 12, 15), "JD", None, "A Person", "VP",
                  "Sent", "a@x.com", "Yes", "7", None, None, None])
    sheet.append(["Recent Inc", "Data Scientist", recent, "JD", None, None, None, None, None, None, None,
                  None, None, None])
    sheet.append(["Example Co", "ML Engineer", "09/12/2025", "duplicate row", None, None, None, None, None, None,
                  None, None, None, None])
    book.save(path)


def plain_profile(pid: str) -> dict:
    """The profile without GOAL4's added columns: these tests check the 14-column mechanics cell by cell."""
    return {k: v for k, v in REAL_PROFILE(pid).items() if k not in ("tracker_add_columns", "tracker_template")}


REAL_PROFILE = config.profile


class TrackerRoundTripTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.conn = db.connect(self.dir / "t.db")
        self.patch = mock.patch.object(config, "profile", plain_profile)
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.conn.close()
        self.tmp.cleanup()

    def roundtrip(self, source: Path) -> tuple[dict, list, list]:
        first = tracker.import_tracker(source, "alex", conn=self.conn)
        out = tracker.export_tracker(self.dir / "out.xlsx", "alex", template=source, conn=self.conn)
        return first, read_cells(source), read_cells(out)

    def test_synthetic_roundtrip_is_lossless_and_idempotent(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        first, original, exported = self.roundtrip(source)
        self.assertEqual(first["inserted"], 4)
        self.assertEqual(original, exported)
        self.assertTrue(all(same_cell(x, y) for row_a, row_b in zip(original, exported) for x, y in zip(row_a, row_b)))
        self.assertIsInstance(exported[1][2], datetime)      # a datetime cell stays a datetime, strings stay strings
        self.assertIsInstance(exported[0][2], str)
        again = tracker.import_tracker(source, "alex", conn=self.conn)
        self.assertEqual((again["inserted"], again["updated"]), (0, 4))
        self.assertEqual(self.conn.execute("SELECT count(*) FROM applications").fetchone()[0], 4)

    def test_status_followups_and_edits(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        apps = {r["company"]: dict(r) for r in self.conn.execute("SELECT * FROM applications")}
        self.assertEqual(apps["Example Co"]["followups_suppressed"], 1)     # 2025 rows never nag
        due = tracker.followups_due(self.conn, "alex")
        self.assertEqual([d["company"] for d in due], ["Recent Inc"])
        self.assertEqual(due[0]["due_1"], (date.today() - timedelta(days=1)).isoformat())
        recent_id = apps["Recent Inc"]["id"]
        tracker.edit_field(self.conn, recent_id, "follow_up_1", "=HYPERLINK(\"x\")")
        tracker.set_status(self.conn, recent_id, "screen")
        out = tracker.export_tracker(self.dir / "out.xlsx", "alex", template=source, conn=self.conn)
        sheet = openpyxl.load_workbook(out).worksheets[0]
        cell = sheet.cell(row=4, column=12)
        self.assertEqual((cell.value, cell.data_type), ("=HYPERLINK(\"x\")", "s"))
        events = [r[0] for r in self.conn.execute("SELECT to_status FROM application_events WHERE app_id=? "
                                                  "ORDER BY id", (recent_id,))]
        self.assertEqual(events, ["applied", "screen"])
        with self.assertRaises(ValueError):
            tracker.set_status(self.conn, recent_id, "hired")

    def test_app_rows_and_pipeline_status_survive_export_import(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        with db.tx(self.conn):                                          # a row created from the web UI / tailor
            self.conn.execute("INSERT INTO applications(profile, job_uid, company, role_title, jd, status) "
                              "VALUES ('alex', 'lever:acme:9', 'Acme', 'ML Engineer', 'https://x/9', 'tailored')")
        recent = self.conn.execute("SELECT id FROM applications WHERE company='Recent Inc'").fetchone()[0]
        tracker.set_status(self.conn, recent, "interview")
        first = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        col = tracker.FIELDS.index("status_ar")
        cells = read_cells(first)
        self.assertEqual({r[0]: r[col] for r in cells}["Recent Inc"], "Interview")
        self.assertEqual({r[0]: r[col] for r in cells}["Acme"], "Tailored")      # as people type it
        self.assertIsNone({r[0]: r[col] for r in cells}["Other AI"])       # untouched rows stay as typed
        again = tracker.import_tracker(first, "alex", conn=self.conn)
        self.assertEqual(again["inserted"], 0)                          # no duplicate for the UI row
        self.assertEqual(self.conn.execute("SELECT count(*) FROM applications").fetchone()[0], 5)
        self.assertEqual(self.conn.execute("SELECT job_uid FROM applications WHERE company='Acme'").fetchone()[0],
                         "lever:acme:9")
        second = tracker.export_tracker(self.dir / "b.xlsx", "alex", template=first, conn=self.conn)
        self.assertEqual(read_cells(first), read_cells(second))         # stable
        book = openpyxl.load_workbook(second)                           # the user marks a rejection in Excel
        sheet = book.worksheets[0]
        for row in sheet.iter_rows(min_row=2):
            if row[0].value == "Other AI":
                row[col].value = "rejected"
        book.save(second)
        tracker.import_tracker(second, "alex", conn=self.conn)
        self.assertEqual(self.conn.execute("SELECT status FROM applications WHERE company='Other AI'").fetchone()[0],
                         "rejected")

    def test_verifier_m5_regressions(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        book = openpyxl.load_workbook(source)
        book.worksheets[0].cell(row=2, column=4).value = "=HYPERLINK(\"http://evil\",\"jd\")"
        book.worksheets[0].cell(row=2, column=4).data_type = "s"
        book.save(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        out = tracker.export_tracker(self.dir / "out.xlsx", "alex", template=source, conn=self.conn)
        self.assertEqual(openpyxl.load_workbook(out).worksheets[0].cell(row=2, column=4).data_type, "s")
        app_id = self.conn.execute("SELECT id FROM applications WHERE company='Other AI'").fetchone()[0]
        # set_status is atomic: an invalid event source rolls the status change back
        with self.assertRaises(Exception):
            tracker.set_status(self.conn, app_id, "screen", source="bogus")
        self.assertEqual(self.conn.execute("SELECT status FROM applications WHERE id=?", (app_id,)).fetchone()[0],
                         "applied")
        # re-marking applied keeps the original applied date
        tracker.set_status(self.conn, app_id, "applied")
        self.assertEqual(self.conn.execute("SELECT applied_on FROM applications WHERE id=?", (app_id,)).fetchone()[0],
                         "2025-12-15")
        # an edited company exported and re-imported does not duplicate the row
        tracker.edit_field(self.conn, app_id, "company", "Other AI Inc")
        out = tracker.export_tracker(self.dir / "out2.xlsx", "alex", template=source, conn=self.conn)
        again = tracker.import_tracker(out, "alex", conn=self.conn)
        self.assertEqual(again["inserted"], 0)
        with self.assertRaises(ValueError):
            tracker.edit_field(self.conn, 999, "company", "x")

    def test_sorted_sheet_with_edited_date_keeps_every_row(self):
        """Release review: the user sorts the exported sheet in Excel and fixes a date; rows keep their identity."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        out = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        book = openpyxl.load_workbook(out)
        sheet = book.worksheets[0]
        rows = [([c.value for c in r], r[0].comment) for r in sheet.iter_rows(min_row=2, max_col=len(tracker.COLUMNS))
                if any(c.value not in (None, "") for c in r)]
        rows.reverse()                                                  # Excel sort: notes move with their cells
        for i, (values, note) in enumerate(rows):
            for col, value in enumerate(values):
                sheet.cell(row=2 + i, column=col + 1, value=value)
            sheet.cell(row=2 + i, column=1).comment = copy.copy(note) if note else None
        edited = next(i for i, (v, _) in enumerate(rows) if v[0] == "Recent Inc")
        sheet.cell(row=2 + edited, column=3, value="01/02/2026")
        book.save(out)
        expected = read_cells(out)
        again = tracker.import_tracker(out, "alex", conn=self.conn)
        self.assertEqual((again["inserted"], again["conflicts"]), (0, []))
        self.assertEqual(self.conn.execute("SELECT count(*) FROM applications").fetchone()[0], 4)
        self.assertEqual(self.conn.execute("SELECT applied_on FROM applications WHERE company='Recent Inc'").fetchone()[0],
                         "2026-01-02")
        final = tracker.export_tracker(self.dir / "b.xlsx", "alex", template=out, conn=self.conn)
        self.assertEqual(read_cells(final), expected)

    def test_identity_edit_without_note_is_adopted_only_when_unambiguous(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        ids = {r["xlsx_row"]: r["id"] for r in self.conn.execute("SELECT id, xlsx_row FROM applications")}
        book = openpyxl.load_workbook(source)                           # no notes: the user's original file
        book.worksheets[0].cell(row=4, column=3, value="01/02/2026")    # Recent Inc: one candidate
        book.worksheets[0].cell(row=2, column=3, value="09/13/2025")    # Example Co: two identical candidates
        book.save(source)
        again = tracker.import_tracker(source, "alex", conn=self.conn)
        self.assertEqual(self.conn.execute("SELECT applied_on FROM applications WHERE company='Recent Inc'").fetchone()[0],
                         "2026-01-02")
        self.assertEqual(again["inserted"], 0)                          # same sheet row breaks the Example Co tie
        after = {r["xlsx_row"]: r["id"] for r in self.conn.execute("SELECT id, xlsx_row FROM applications")}
        self.assertEqual(after, ids)                                    # the two Example Co rows did not swap

    def test_sheet_value_over_app_edit_is_reported(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        app_id = self.conn.execute("SELECT id FROM applications WHERE company='Other AI'").fetchone()[0]
        tracker.edit_field(self.conn, app_id, "target_name", "From App")
        again = tracker.import_tracker(source, "alex", conn=self.conn)
        self.assertEqual(again["kept_app_values"],
                         [{"row": 3, "field": "target_name", "sheet": "A Person", "app": "From App"}])

    def _sort_reverse(self, path: Path) -> None:
        """What Excel's sort does: values, hyperlinks and notes move with their row."""
        book = openpyxl.load_workbook(path)
        sheet = book.worksheets[0]
        rows = [[(c.value, copy.copy(c.hyperlink) if c.hyperlink else None, copy.copy(c.comment) if c.comment else None)
                 for c in r] for r in sheet.iter_rows(min_row=2, max_col=len(tracker.COLUMNS))
                if any(c.value not in (None, "") for c in r)]
        for i, row in enumerate(reversed(rows)):
            for col, (value, link, note) in enumerate(row):
                cell = sheet.cell(row=2 + i, column=col + 1)
                cell.hyperlink = None
                cell.value = value
                if link is not None:
                    cell.hyperlink = link.target
                cell.comment = note
        book.save(path)

    @staticmethod
    def _links(path: Path) -> dict:
        sheet = openpyxl.load_workbook(path).worksheets[0]
        return {(r[0].value, r[3].value, c.column): (c.value, c.hyperlink.target if c.hyperlink else None)
                for r in sheet.iter_rows(min_row=2, max_col=len(tracker.COLUMNS)) if r[0].value
                for c in r if c.hyperlink or c.column == 5}

    def test_hyperlinks_move_with_their_rows(self):
        """Verifier M9: a sorted sheet must not leave links behind on the template's old rows."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        book = openpyxl.load_workbook(source)
        sheet = book.worksheets[0]
        sheet.cell(row=2, column=4).hyperlink = "https://jobs.example/1"           # JD text with a link
        sheet.cell(row=3, column=5, value="resume v2").hyperlink = "https://drive.example/2"
        book.save(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        first = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        expected = self._links(first)
        self._sort_reverse(first)
        self.assertEqual(self._links(first), expected)                              # the sort itself is faithful
        tracker.import_tracker(first, "alex", conn=self.conn)
        second = tracker.export_tracker(self.dir / "b.xlsx", "alex", template=source, conn=self.conn)
        self.assertEqual(self._links(second), expected)
        self.assertEqual(read_cells(second), read_cells(first))

    def test_legacy_rows_keep_template_links(self):
        """Rows imported before links were recorded still export the link at their own sheet row."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        book = openpyxl.load_workbook(source)
        book.worksheets[0].cell(row=2, column=4).hyperlink = "https://jobs.example/1"
        book.save(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        with db.tx(self.conn):
            for app_id, raw in self.conn.execute("SELECT id, raw_json FROM applications").fetchall():
                legacy = [{k: v for k, v in cell.items() if k != "h"} for cell in json.loads(raw)]
                self.conn.execute("UPDATE applications SET raw_json=? WHERE id=?", (json.dumps(legacy), app_id))
        out = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        self.assertEqual(openpyxl.load_workbook(out).worksheets[0].cell(row=2, column=4).hyperlink.target,
                         "https://jobs.example/1")

    def test_stale_export_never_moves_status_back(self):
        """Verifier M9: importing an older export keeps a status changed in the app since; a status the user typed
        in the sheet still wins."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        recent = self.conn.execute("SELECT id FROM applications WHERE company='Recent Inc'").fetchone()[0]
        tracker.set_status(self.conn, recent, "interview")
        old = tracker.export_tracker(self.dir / "old.xlsx", "alex", template=source, conn=self.conn)
        tracker.set_status(self.conn, recent, "offer")
        tracker.import_tracker(old, "alex", conn=self.conn)
        status = lambda: self.conn.execute("SELECT status FROM applications WHERE id=?", (recent,)).fetchone()[0]
        self.assertEqual(status(), "offer")
        new = tracker.export_tracker(self.dir / "new.xlsx", "alex", template=source, conn=self.conn)
        book = openpyxl.load_workbook(new)
        col = tracker.FIELDS.index("status_ar") + 1
        for row in book.worksheets[0].iter_rows(min_row=2):
            if row[0].value == "Recent Inc":
                row[col - 1].value = "rejected"
        book.save(new)
        tracker.import_tracker(new, "alex", conn=self.conn)
        self.assertEqual(status(), "rejected")

    def test_an_older_export_never_moves_status_back_after_a_newer_one(self):
        """Verifier M9 round 2: what THIS file showed decides, not what the latest export showed."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        recent = self.conn.execute("SELECT id FROM applications WHERE company='Recent Inc'").fetchone()[0]
        status = lambda: self.conn.execute("SELECT status FROM applications WHERE id=?", (recent,)).fetchone()[0]
        tracker.set_status(self.conn, recent, "interview")
        older = tracker.export_tracker(self.dir / "e1.xlsx", "alex", template=source, conn=self.conn)
        tracker.set_status(self.conn, recent, "offer")
        tracker.export_tracker(self.dir / "e2.xlsx", "alex", template=source, conn=self.conn)
        result = tracker.import_tracker(older, "alex", conn=self.conn)
        self.assertEqual(status(), "offer")
        self.assertNotIn("status_ar", [k["field"] for k in result["kept_app_values"]])   # untouched cell: no conflict
        book = openpyxl.load_workbook(older)                            # the user edits the old file's status
        col = tracker.FIELDS.index("status_ar")
        for row in book.worksheets[0].iter_rows(min_row=2):
            if row[0].value == "Recent Inc":
                row[col].value = "rejected"
        book.save(older)
        result = tracker.import_tracker(older, "alex", conn=self.conn)
        self.assertEqual(status(), "offer")                             # the app changed since that file: app wins
        self.assertIn({"row": 4, "field": "status_ar", "sheet": "rejected", "app": "offer"}, result["kept_app_values"])

    def test_v7_backfills_what_pre_upgrade_rows_showed(self):
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        with db.tx(self.conn):
            self.conn.execute("UPDATE applications SET sheet_status=NULL, status_ar='Rejected ' WHERE company='Recent Inc'")
            self.conn.execute("UPDATE applications SET sheet_status=NULL WHERE company='Other AI'")
            self.conn.execute(db.MIGRATIONS[7][0])
        shown = dict(self.conn.execute("SELECT company, sheet_status FROM applications WHERE company IN "
                                       "('Recent Inc', 'Other AI')").fetchall())
        self.assertEqual(shown, {"Recent Inc": "rejected", "Other AI": ""})

    def test_note_free_file_never_moves_status_back(self):
        """Verifier M9 round 3: the user's own tracker has no notes; a status the app already left is reported."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        col = tracker.FIELDS.index("status_ar") + 1
        book = openpyxl.load_workbook(source)
        book.worksheets[0].cell(row=4, column=col, value="interview")         # Recent Inc, typed by the user
        book.save(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        recent = self.conn.execute("SELECT id FROM applications WHERE company='Recent Inc'").fetchone()[0]
        status = lambda: self.conn.execute("SELECT status FROM applications WHERE id=?", (recent,)).fetchone()[0]
        self.assertEqual(status(), "interview")
        tracker.set_status(self.conn, recent, "offer")
        tracker.export_tracker(self.dir / "e.xlsx", "alex", template=source, conn=self.conn)
        result = tracker.import_tracker(source, "alex", conn=self.conn)          # the old master file again
        self.assertEqual(status(), "offer")
        self.assertIn({"row": 4, "field": "status_ar", "sheet": "interview", "app": "offer"}, result["kept_app_values"])
        book = openpyxl.load_workbook(source)                                   # a status it never had applies
        book.worksheets[0].cell(row=4, column=col, value="rejected")
        book.save(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        self.assertEqual(status(), "rejected")

    def test_copied_row_never_takes_over_its_source(self):
        """Verifier M9 round 3: Excel copies the note with the row; the row that matches the app keeps it."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        out = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        book = openpyxl.load_workbook(out)
        sheet = book.worksheets[0]
        sheet.insert_rows(2)                                                    # the copy lands above its source
        for col in range(1, len(tracker.COLUMNS) + 1):
            sheet.cell(row=2, column=col, value=sheet.cell(row=4, column=col).value)   # Other AI is now row 4
        sheet.cell(row=2, column=1).comment = copy.copy(sheet.cell(row=4, column=1).comment)
        sheet.cell(row=2, column=1, value="NewCo Robotics")
        sheet.cell(row=2, column=2, value="ML Engineer")
        sheet.cell(row=2, column=3, value="10/01/2026")
        sheet.cell(row=2, column=len(tracker.COLUMNS), value="interview")
        book.save(out)
        before = self.conn.execute("SELECT id, status FROM applications WHERE company='Other AI'").fetchone()
        result = tracker.import_tracker(out, "alex", conn=self.conn)
        after = self.conn.execute("SELECT company, status FROM applications WHERE id=?", (before["id"],)).fetchone()
        self.assertEqual(tuple(after), ("Other AI", before["status"]))
        self.assertEqual(result["inserted"], 1)
        self.assertEqual(self.conn.execute("SELECT status FROM applications WHERE company='NewCo Robotics'")
                         .fetchone()[0], "interview")

    def _insert_copy_above(self, path: Path, company: str, edits: dict, source_edits: dict | None = None) -> None:
        """Excel 'Insert Copied Cells': the copy takes the source's row, the source moves down one."""
        book = openpyxl.load_workbook(path)
        sheet = book.worksheets[0]
        src = next(r[0].row for r in sheet.iter_rows(min_row=2) if r[0].value == company)
        sheet.insert_rows(src)
        for col in range(1, len(tracker.COLUMNS) + 1):
            sheet.cell(row=src, column=col, value=sheet.cell(row=src + 1, column=col).value)
        sheet.cell(row=src, column=1).comment = copy.copy(sheet.cell(row=src + 1, column=1).comment)
        for col, value in edits.items():
            sheet.cell(row=src, column=col, value=value)
        for col, value in (source_edits or {}).items():
            sheet.cell(row=src + 1, column=col, value=value)
        book.save(path)

    def test_tied_copy_never_silently_takes_over(self):
        """Verifier M9 round 4: copy and source tie on identity; the source keeps the app and the tie is reported."""
        status_col = len(tracker.COLUMNS)
        cases = [({1: "Ramp"}, {2: "Senior AI Engineer – Agents"}, "Other AI"),        # company kept wins
                 ({status_col: "offer"}, None, "Other AI"),                             # unchanged row wins
                 ({6: "Second Contact"}, None, "Other AI")]
        for edits, source_edits, company in cases:
            with self.subTest(edits=edits):
                self.tearDown()
                self.setUp()
                source = self.dir / "in.xlsx"
                make_fixture(source)
                tracker.import_tracker(source, "alex", conn=self.conn)
                app = dict(self.conn.execute("SELECT * FROM applications WHERE company=?", (company,)).fetchone())
                out = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
                self._insert_copy_above(out, company, edits, source_edits)
                result = tracker.import_tracker(out, "alex", conn=self.conn)
                after = dict(self.conn.execute("SELECT * FROM applications WHERE id=?", (app["id"],)).fetchone())
                self.assertEqual((after["company"], after["status"], after["target_name"]),
                                 (app["company"], app["status"], app["target_name"]))
                self.assertEqual(result["inserted"], 1)
                self.assertTrue(result["conflicts"], "a tie is always reported")

    def test_tie_breaks_use_closest_company_and_the_files_own_status(self):
        """Verifier M9 round 5: a company spelling fix on the source beats a new company on the copy; and cells are
        compared with what that file showed, not with the app's newer status."""
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        app_id = self.conn.execute("SELECT id FROM applications WHERE company='Other AI'").fetchone()[0]
        out = tracker.export_tracker(self.dir / "a.xlsx", "alex", template=source, conn=self.conn)
        self._insert_copy_above(out, "Other AI", {1: "Ramp"}, {1: "Other A.I"})
        result = tracker.import_tracker(out, "alex", conn=self.conn)
        self.assertEqual(self.conn.execute("SELECT company FROM applications WHERE id=?", (app_id,)).fetchone()[0],
                         "Other A.I")
        self.assertTrue(result["conflicts"])

        self.tearDown()
        self.setUp()
        source = self.dir / "in.xlsx"
        make_fixture(source)
        tracker.import_tracker(source, "alex", conn=self.conn)
        app_id = self.conn.execute("SELECT id FROM applications WHERE company='Other AI'").fetchone()[0]
        old = tracker.export_tracker(self.dir / "old.xlsx", "alex", template=source, conn=self.conn)
        tracker.set_status(self.conn, app_id, "interview")                     # the app moved on after that file
        self._insert_copy_above(old, "Other AI", {len(tracker.COLUMNS): "screen", 6: "New Person"})
        result = tracker.import_tracker(old, "alex", conn=self.conn)
        after = self.conn.execute("SELECT status, target_name FROM applications WHERE id=?", (app_id,)).fetchone()
        self.assertEqual(tuple(after), ("interview", "A Person"))
        self.assertEqual(self.conn.execute("SELECT target_name FROM applications WHERE id<>? AND company='Other AI'",
                                           (app_id,)).fetchone()[0], "New Person")     # the copy's data is kept

    @unittest.skipUnless(REAL.exists(), "real tracker not present (data/ is gitignored)")
    def test_real_tracker_roundtrip_374_rows(self):
        first, original, exported = self.roundtrip(REAL)
        self.assertEqual(len(original), 374)
        self.assertEqual(first["rows"], 374)
        mismatches = [(r, c) for r, (a, b) in enumerate(zip(original, exported))
                      for c in range(len(tracker.COLUMNS)) if not same_cell(a[c], b[c])]
        self.assertEqual(len(exported), 374)
        self.assertEqual(mismatches, [])
        again = tracker.import_tracker(REAL, "alex", conn=self.conn)
        self.assertEqual(again["inserted"], 0)


if __name__ == "__main__":
    unittest.main()
