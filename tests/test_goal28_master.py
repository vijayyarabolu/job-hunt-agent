"""GOAL28 §1: `job-hunt-agent baseline master` normalizes ONE master resume (links show their address, dates on one right tab,
nothing justified), archives the old bases (nothing deleted) and rebuilds the baseline from it."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn

from job_hunt_agent.resume import baseline, master

PAD = " " * 40
DAY = date(2026, 10, 18)


def add_link(paragraph, url: str, text: str) -> None:
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = paragraph._p.makeelement(qn("w:hyperlink"), {qn("r:id"): rid})
    run = link.makeelement(qn("w:r"), {})
    t = run.makeelement(qn("w:t"), {})
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def synthetic_master(path: Path) -> None:
    """A master as typed in Word: header links say 'LinkedIn', the dates sit after typed spaces split over runs (the
    padding starts exactly at a run boundary for the second job), the summary is justified."""
    doc = docx.Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph("AI Engineer")
    contact = doc.add_paragraph("555-000-0000 | jane@example.com | ")
    add_link(contact, "https://www.linkedin.com/in/jane-doe/", "LinkedIn")
    doc.add_paragraph("Summary")
    doc.add_paragraph("AI engineer with 4 years building LLM systems.").alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph("Skills")
    doc.add_paragraph("Programming: Python, SQL")
    doc.add_paragraph("Work Experience")
    first = doc.add_paragraph()
    first.add_run("New Co, USA" + PAD[:10])
    first.add_run(PAD[10:] + "Jul 2025 – Present")
    doc.add_paragraph("AI Engineer")
    doc.add_paragraph("Built an LLM system in Python improving latency 30% for 200 users.", style="List Bullet")
    second = doc.add_paragraph()
    second.add_run("Old Inc, India")
    second.add_run(PAD + "Jan 2020 – Jun 2023")
    doc.add_paragraph("Data Engineer")
    doc.add_paragraph("Built Spark pipelines in Python cutting run time 40% for the risk team.", style="List Bullet")
    doc.add_paragraph("Education")
    doc.add_paragraph("MS in Computer Science | Example University" + PAD + "May 2025")
    doc.save(path)


def tab_stops(paragraph) -> list[tuple[str, str]]:
    return [(t.get(qn("w:val")), t.get(qn("w:pos"))) for t in paragraph._p.iter(qn("w:tab")) if t.get(qn("w:val"))]


class LinkTextTest(unittest.TestCase):
    def test_the_address_is_what_a_reader_can_type(self):
        self.assertEqual(master.link_text("https://www.linkedin.com/in/x/"), "linkedin.com/in/x")
        self.assertEqual(master.link_text("http://github.com/y"), "github.com/y")


class NormalizeTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        synthetic_master(self.dir / "m.docx")
        self.out = master.normalize(self.dir / "m.docx", self.dir / "n.docx")
        self.doc = docx.Document(str(self.dir / "n.docx"))

    def tearDown(self):
        self.tmp.cleanup()

    def test_header_links_show_their_address_and_stay_links(self):
        self.assertEqual(self.out["links"], ["linkedin.com/in/jane-doe"])
        self.assertIn("linkedin.com/in/jane-doe", "".join(t.text for t in self.doc.element.body.iter(qn("w:t"))))
        self.assertEqual(len(list(self.doc.element.body.iter(qn("w:hyperlink")))), 1)

    def test_every_dated_line_gets_one_tab_and_one_right_stop(self):
        self.assertEqual(self.out["dates_tabbed"], ["New Co, USA", "Old Inc, India",
                                                    "MS in Computer Science | Example University"])
        dated = [p for p in self.doc.paragraphs if "\t" in p.text]
        self.assertEqual([p.text for p in dated], ["New Co, USA\tJul 2025 – Present", "Old Inc, India\tJan 2020 – Jun 2023",
                                                   "MS in Computer Science | Example University\tMay 2025"])
        for p in dated:
            stops = tab_stops(p)
            self.assertEqual(len(stops), 1)
            self.assertEqual(stops[0][0], "right")

    def test_the_right_stop_sits_just_inside_the_text_edge(self):
        section = self.doc.sections[0]
        edge = (section.page_width - section.left_margin - section.right_margin) // master.EMU_PER_TWIP
        p = next(p for p in self.doc.paragraphs if p.text.startswith("Old Inc"))
        self.assertEqual(int(tab_stops(p)[0][1]), edge - master.RIGHT_TAB_SLACK)

    def test_nothing_is_justified(self):
        self.assertEqual(self.out["unjustified"], 1)
        self.assertFalse([jc for jc in self.doc.element.body.iter(qn("w:jc")) if jc.get(qn("w:val")) == "both"])

    def test_the_parse_is_unchanged_apart_from_layout(self):
        before, after = baseline.parse(self.dir / "m.docx"), baseline.parse(self.dir / "n.docx")
        self.assertEqual(after["experience"], before["experience"])
        self.assertEqual(after["summary"], before["summary"])
        self.assertEqual(after["skills"], before["skills"])


class InstallTest(unittest.TestCase):
    def test_old_bases_are_archived_and_one_base_remains(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            synthetic_master(tmp / "m.docx")
            profile = tmp / "p"
            (profile / "resumes").mkdir(parents=True)
            for name in ("ml_engineer.docx", "Old_Resume.docx", "Old_Resume.pdf"):
                (profile / "resumes" / name).write_bytes(b"old")
            (tmp / "m.pdf").write_bytes(b"%PDF")
            got = master.install(tmp / "m.docx", profile, "Jane_Master", tmp / "m.pdf", DAY)
            archive = profile / "resumes" / "archive" / "2026-10-18"
            self.assertEqual(sorted(got["archived"]), ["Old_Resume.docx", "Old_Resume.pdf", "ml_engineer.docx"])
            self.assertEqual(sorted(p.name for p in archive.iterdir()), sorted(got["archived"]))
            self.assertEqual(sorted(p.name for p in (profile / "resumes").iterdir() if p.is_file()),
                             ["Jane_Master.docx", "Jane_Master.pdf"])
            self.assertTrue((profile / "baseline.json").exists())
            self.assertEqual(got["resume"]["headline"], "AI Engineer")

    def test_a_second_archive_the_same_day_keeps_both(self):
        with tempfile.TemporaryDirectory() as tmp:
            resumes = Path(tmp)
            (resumes / "a.docx").write_bytes(b"1")
            master.archive_old(resumes, DAY)
            (resumes / "a.docx").write_bytes(b"2")
            master.archive_old(resumes, DAY)
            names = sorted(p.name for p in (resumes / "archive" / "2026-10-18").iterdir())
        self.assertEqual(names, ["a.1.docx", "a.docx"])

    def test_an_unparseable_master_changes_nothing(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            doc = docx.Document()
            doc.add_paragraph("just a letter")
            doc.save(tmp / "bad.docx")
            (tmp / "p" / "resumes").mkdir(parents=True)
            (tmp / "p" / "resumes" / "old.docx").write_bytes(b"old")
            with self.assertRaises(ValueError):
                master.install(tmp / "bad.docx", tmp / "p", "X", day=DAY)
            self.assertEqual([p.name for p in (tmp / "p" / "resumes").iterdir()], ["old.docx"])
            self.assertEqual([p.name for p in (tmp / "p").iterdir()], ["resumes"])


if __name__ == "__main__":
    unittest.main()
