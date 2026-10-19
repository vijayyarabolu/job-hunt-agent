"""`baseline import`: generic .docx -> baseline.json parsing (any number of jobs/bullets, no hardcoding)."""
import json
import tempfile
import unittest
from pathlib import Path

import docx

from job_hunt_agent import config
from job_hunt_agent.resume import baseline

TEMPLATE = config.PROFILES / "alex" / "template.docx"


def synthetic_resume(path: Path, jobs: list[tuple[str, str, str, int]]) -> None:
    doc = docx.Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph("Machine Learning Engineer")
    doc.add_paragraph("555-000-0000 | jane@example.com | Austin, TX")
    doc.add_paragraph("Summary\t")
    doc.add_paragraph("ML engineer with 5+ years building ranking systems.")
    doc.add_paragraph("Skills\t ")
    doc.add_paragraph("Cloud: AWS (S3, Glue), GCP")
    doc.add_paragraph("Programming: Python, SQL")
    doc.add_paragraph("Education\t ")
    doc.add_paragraph("MS in Computer Science | Example University, Austin, TX")
    doc.add_paragraph("Work Experience\t\t ")
    for company, dates, title, n_bullets in jobs:
        doc.add_paragraph(f"{company}\t{dates}")
        doc.add_paragraph(title)
        for i in range(n_bullets):
            doc.add_paragraph(f"Built system {i} using Python improving latency {i + 10}%.", style="List Bullet")
    doc.save(path)


class BaselineTest(unittest.TestCase):
    def test_any_number_of_jobs_and_bullets(self):
        jobs = [("New Co, USA", "Mar 2026 – Present", "ML Engineer", 3),
                ("Example Corp, USA", "Jul 2023 – Feb 2026", "Data Scientist", 6),
                ("Old Inc, India", "2019 - 2023", "Analyst", 1)]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.docx"
            synthetic_resume(path, jobs)
            resume = baseline.import_baseline(path, Path(tmp) / "profile")
            saved = json.loads((Path(tmp) / "profile" / "baseline.json").read_text())
            known = (Path(tmp) / "profile" / "known_skills.txt").read_text()
        self.assertEqual(saved, {k: v for k, v in resume.items() if k != "_import"})
        self.assertEqual([(j["company"], j["dates"], j["title"], len(j["bullets"])) for j in resume["experience"]], jobs)
        self.assertEqual(resume["headline"], "Machine Learning Engineer")
        self.assertEqual(resume["skills"][0], {"category": "Cloud", "items": ["AWS (S3, Glue)", "GCP"]})
        self.assertEqual(resume["education"], [{"degree": "MS in Computer Science",
                                                "school": "Example University, Austin, TX"}])
        for term in ("AWS", "S3", "Glue", "GCP", "Python"):
            self.assertIn(f"\n{term}\n", known)

    def test_existing_known_skills_are_kept_and_new_baseline_terms_added(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.docx"
            synthetic_resume(path, [("A, USA", "2020 - 2021", "Engineer", 2)])
            profile = Path(tmp) / "profile"
            profile.mkdir()
            (profile / "known_skills.txt").write_text("Rust\naws\n")
            result = baseline.import_baseline(path, profile)
            lines = (profile / "known_skills.txt").read_text().splitlines()
        self.assertEqual(lines[:2], ["Rust", "aws"])
        self.assertNotIn("AWS", lines)                      # case-insensitive union
        self.assertIn("GCP", lines)
        self.assertIn("GCP", result["_import"]["known_added"])

    def test_import_installs_template_and_keeps_the_old_one(self):
        with tempfile.TemporaryDirectory() as tmp:
            first, second = Path(tmp) / "a.docx", Path(tmp) / "b.docx"
            synthetic_resume(first, [("A, USA", "2020 - 2021", "Engineer", 2)])
            synthetic_resume(second, [("B, USA", "Jan 2022 – Present", "Engineer", 3),
                                      ("A, USA", "2020 - 2021", "Engineer", 2)])
            profile = Path(tmp) / "profile"
            baseline.import_baseline(first, profile)
            result = baseline.import_baseline(second, profile)
            self.assertEqual((profile / "template.docx").read_bytes(), second.read_bytes())
            backup = Path(result["_import"]["template_backup"])
            self.assertEqual(backup.read_bytes(), first.read_bytes())
            self.assertEqual(len(json.loads((profile / "baseline.json").read_text())["experience"]), 2)

    def test_unparseable_resume_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "r.docx"
            docx.Document().save(path)
            with self.assertRaises(ValueError):
                baseline.import_baseline(path, Path(tmp) / "profile")

    @unittest.skipUnless(TEMPLATE.exists(), "personal template not present")
    def test_real_template_structure(self):
        resume = baseline.parse(TEMPLATE)
        self.assertEqual(baseline.validate(resume), [])
        self.assertTrue(all(j["title"] and j["dates"] and j["bullets"] for j in resume["experience"]))
        stored = TEMPLATE.parent / "baseline.json"
        if stored.exists():
            self.assertEqual(resume, json.loads(stored.read_text()))


if __name__ == "__main__":
    unittest.main()
