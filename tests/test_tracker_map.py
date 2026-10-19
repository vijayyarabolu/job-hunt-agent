"""Mapped tracker import: a workbook with its own layout lands in the normal applications table."""
import json
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest import mock

import openpyxl
import yaml

from job_hunt_agent import config, db, tracker, tracker_map

MAPPING = {
    "extras": ["Address", "Team Name"],
    "swap_date_cells": True,
    "date_fixes": {"09/31/2026": "09/30/2026"},
    "sheets": [
        {"name": "Main", "columns": {"Company": "Company Name", "Role": "Role Title", "Date": "Date Applied",
                                     "Status": "Status(A/R)", "Address": "Address"}},
        {"name": "Second", "columns": {"Company": "Company Name", "Role": "Role Title", "Date": "Date Applied",
                                       "Status": "Status(A/R)", "Addresss": "Address", "Team Name": "Team Name"},
         "fill_if_empty": {"Column 1": "Status(A/R)"}},
    ],
    "vendors": {"sheet": "Vendors", "columns": {"Vendor Companies": "company", "Mail": "email"}},
}


def extra(raw_json: str, header: str):
    """An extra column the way the Sheet tab reads it: the tagged cell carrying 'col' == header."""
    return next(tracker.decode(c) for c in json.loads(raw_json) if c.get("col") == header)


def make_workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    main = book.active
    main.title = "Main"
    main.append(["Company", "Role", "Date", "Status", "Address"])
    main.append(["Acme", "Data Analyst", datetime(2026, 5, 9), None, "Austin, TX"])      # typed 09/05 -> Sep 5
    main.append(["Beta", "BI Analyst", "09/31/2026", "Rejected", None])                 # approved fix
    main.append(["Beta", "BI Analyst", "09/31/2026", None, None])                       # repeat inside one sheet
    main.append([None, None, None, None, None])
    main.append([None, None, datetime(2026, 2, 10), None, None])                         # pre-filled date only
    second = book.create_sheet("Second")
    second.append(["Company", "Role", "Date", "Status", "Addresss", "Team Name", "Column 1"])
    second.append(["Acme", "Data Analyst", datetime(2026, 5, 9), None, None, None, None])  # same app as Main
    second.append(["Gamma", "Analyst", "08/01/2026", None, "Remote", "Payments", "R"])
    vendors = book.create_sheet("Vendors")
    vendors.append(["Vendor Companies", "Mail"])
    vendors.append(["StaffCo", "a@staff.example"])
    vendors.append([None, None])
    book.save(path)


class TrackerMapTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.xlsx = self.dir / "theirs.xlsx"
        make_workbook(self.xlsx)
        self.map = self.dir / "map.yaml"
        self.map.write_text(yaml.safe_dump(MAPPING))
        (self.dir / "profiles" / "p").mkdir(parents=True)
        self.patch = mock.patch.object(config, "PROFILES", self.dir / "profiles")
        self.patch.start()
        self.conn = db.connect(self.dir / "t.db")

    def tearDown(self):
        self.conn.close()
        self.patch.stop()
        self.tmp.cleanup()

    def app(self, company: str) -> dict:
        return dict(self.conn.execute("SELECT * FROM applications WHERE profile='p' AND company=? ORDER BY id",
                                      (company,)).fetchone())

    def count(self, where: str = "1") -> int:
        return self.conn.execute(f"SELECT count(*) FROM applications WHERE {where}").fetchone()[0]

    def test_dry_run_counts_and_writes_nothing(self):
        result = tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn, dry=True)
        self.assertEqual(result["rows"], 4)
        self.assertEqual(result["per_sheet"], {"Main": 3, "Second": 2})
        self.assertEqual(result["cross_sheet_duplicates"], 1)
        self.assertEqual(result["within_sheet_repeats"], 1)
        self.assertEqual(result["vendors"], 1)
        self.assertEqual(result["unparsed_dates"], [])
        self.assertFalse((self.dir / "profiles" / "p" / "vendors.json").exists())
        self.assertEqual(self.count(), 0)

    def test_import_maps_dates_status_extras_and_vendors(self):
        result = tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn)
        self.assertEqual(result["inserted"], 4)
        self.assertEqual(self.app("Acme")["applied_on"], "2026-09-05")            # swapped date cell
        self.assertEqual(self.count("company='Beta' AND applied_on='2026-09-30'"), 2)  # fix + kept repeat
        gamma = self.app("Gamma")
        self.assertEqual(gamma["status_ar"], "R")                                  # Column 1 filled empty Status
        self.assertEqual(extra(gamma["raw_json"], "Address"), "Remote")
        self.assertEqual(extra(gamma["raw_json"], "Team Name"), "Payments")
        self.assertEqual(extra(self.app("Acme")["raw_json"], "Address"), "Austin, TX")
        vendors = json.loads((self.dir / "profiles" / "p" / "vendors.json").read_text())
        self.assertEqual(vendors, [{"company": "StaffCo", "email": "a@staff.example"}])

    def test_reimport_is_idempotent(self):
        tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn)
        again = tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn)
        self.assertEqual((again["inserted"], again["updated"]), (0, 4))
        self.assertEqual(self.count(), 4)

    def test_unknown_target_column_is_rejected(self):
        self.map.write_text(yaml.safe_dump({**MAPPING, "sheets": [{"name": "Main", "columns": {"Company": "Firm"}}]}))
        with self.assertRaisesRegex(ValueError, "unknown column 'Firm'"):
            tracker_map.load_mapping(self.map)

    def test_missing_source_column_is_rejected(self):
        bad = {**MAPPING, "sheets": [{"name": "Main", "columns": {"Employer": "Company Name"}}]}
        self.map.write_text(yaml.safe_dump(bad))
        with self.assertRaisesRegex(ValueError, "no column"):
            tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn, dry=True)

    def test_day_after_twelve_is_never_swapped(self):
        mapping = tracker_map.load_mapping(self.map)
        self.assertEqual(tracker_map.fix_date(datetime(2026, 5, 20), mapping), datetime(2026, 5, 20))
        self.assertEqual(tracker_map.fix_date("09/31/2026", mapping), "09/30/2026")
        self.assertEqual(tracker_map.fix_date("08/01/2026", mapping), "08/01/2026")


if __name__ == "__main__":
    unittest.main()


def by_header(sheet) -> list[dict]:
    head = [c.value for c in sheet[1]]
    return [dict(zip(head, (c.value for c in row))) for row in sheet.iter_rows(min_row=2) if row[0].value is not None
            or any(c.value is not None for c in row)]


class ExportMappedTest(TrackerMapTest):
    """the tracker back in the workbook's own layout; importing it again changes nothing."""

    def snapshot(self) -> dict:
        cols = ", ".join(("company", "role_title", "date_applied", "applied_on", "status", "status_ar", "address",
                          "team_name", "target_name", "email_address"))
        return {r["id"]: tuple(r) for r in self.conn.execute(f"SELECT id, {cols} FROM applications ORDER BY id")}

    def test_export_round_trips_rows_dates_extras_and_vendors(self):
        tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn)
        before = self.snapshot()
        out = tracker_map.export_mapped(self.conn, "p", self.map, self.dir / "out.xlsx")
        book = openpyxl.load_workbook(out)
        self.assertEqual([s.title for s in book.worksheets], ["Main", "Second", "Vendors", "job-hunt-agent app fields"])
        second = by_header(book["Second"])[0]                           # Gamma came from its own sheet
        self.assertEqual({k: second[k] for k in ("Company", "Role", "Date", "Status", "Addresss", "Team Name")},
                         {"Company": "Gamma", "Role": "Analyst", "Date": "08/01/2026", "Status": "R",
                          "Addresss": "Remote", "Team Name": "Payments"})
        self.assertEqual(by_header(book["Main"])[0]["Date"], datetime(2026, 9, 5))      # the corrected date
        self.assertIn("job-hunt-agent #", book["Main"]["A2"].comment.text)
        self.assertEqual(by_header(book["Vendors"]), [{"Vendor Companies": "StaffCo", "Mail": "a@staff.example"}])
        again = tracker_map.import_mapped(out, "p", self.map, self.conn)
        self.assertEqual((again["inserted"], again["updated"], again["conflicts"]), (0, 4, []))
        self.assertEqual(self.snapshot(), before)                     # no swap of the corrected date, nothing lost

    def test_edited_extra_is_exported_from_the_app(self):
        from job_hunt_agent import trackapi
        tracker_map.import_mapped(self.xlsx, "p", self.map, self.conn)
        gamma = self.app("Gamma")["id"]
        trackapi.update(self.conn, "p", gamma, {"team_name": "Risk"})
        book = openpyxl.load_workbook(tracker_map.export_mapped(self.conn, "p", self.map, self.dir / "out.xlsx"))
        self.assertEqual(by_header(book["Second"])[0]["Team Name"], "Risk")
