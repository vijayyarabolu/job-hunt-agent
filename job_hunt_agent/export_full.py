"""Full tracker export: every field of every application of a profile, one row each, plus a Contacts sheet
(several per application) and a Vendors sheet. A read-only view for backups and analysis; the per-person exports in
their own layouts (tracker.export_tracker / tracker_map.export_mapped) are the ones that import back."""
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from . import config, trackapi, vendors

APP_COLUMNS = (("id", "Row ID"), ("company", "Company Name"), ("role_title", "Role Title"),
               ("date_applied", "Date Applied"), ("applied_on", "Applied On (ISO)"), ("status", "Pipeline Status"),
               ("status_ar", "Status(A/R)"), ("source", "Source"), ("family", "Role Family"), ("jd", "JD"),
               ("resume", "Resume"), ("job_url", "Job URL"), ("apply_url", "Apply URL"), ("location", "Location"),
               ("workplace", "Workplace"), ("salary", "Salary"), ("employment_label", "Employment Type"),
               ("address", "Address"), ("team_name", "Team Name"), ("target_name", "Target Name"),
               ("target_role", "Target Role"), ("linkedin_request", "LinkedIn Request"),
               ("email_address", "Email Address"), ("email_sent", "Email Sent"), ("gemini_chat", "Gemini chat number"),
               ("follow_up_1", "Follow-Up 1(3days)"), ("fu1_due", "Follow-up 1 due"),
               ("follow_up_2", "Follow-Up 2(5-7days)"), ("fu2_due", "Follow-up 2 due"), ("notes", "Notes"),
               ("job_uid", "Job UID"), ("posted_at", "Posted"), ("updated_at", "Last Updated"))
CONTACT_COLUMNS = (("app_id", "Row ID"), ("company", "Company Name"), ("role_title", "Role Title"), ("name", "Name"),
                   ("role", "Role"), ("email", "Email"), ("linkedin", "LinkedIn"),
                   ("linkedin_sent", "LinkedIn sent"), ("email_sent", "Email sent"), ("replied", "Replied"),
                   ("notes", "Notes"))
VENDOR_COLUMNS = (("company", "Vendor Companies"), ("target_name", "Target name"), ("target_role", "Target role"),
                  ("linkedin", "Linkedin connection"), ("email", "Mail"), ("email_sent", "Mail Sent"),
                  ("notes", "Notes"))
_FORMULA = ("=", "+", "-", "@")


def _sheet(book, title: str, columns: tuple, rows: list[dict]) -> None:
    sheet = book.create_sheet(title)
    sheet.append([label for _, label in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in rows:
        values = [row.get(key) for key, _ in columns]
        sheet.append(values)
        for cell, value in zip(sheet[sheet.max_row], values):
            if isinstance(value, str) and value.startswith(_FORMULA):    # text, never a formula
                cell.data_type = "s"
                cell.quotePrefix = True


def all_rows(conn, pid: str) -> list[dict]:
    """Every tracker row of the profile (trackapi views, paged), by row id."""
    page = trackapi.rows(conn, pid, limit=trackapi.MAX_PAGE)
    apps = page["rows"]
    while len(apps) < page["total"]:
        page = trackapi.rows(conn, pid, limit=trackapi.MAX_PAGE, offset=len(apps))
        apps += page["rows"]
    return sorted(apps, key=lambda a: a["id"])


def export_full(conn, pid: str, out: Path | None = None) -> Path:
    out = Path(out or config.DATA / "export" / f"Tracker_full_{pid}_{date.today():%Y%m%d}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    apps = all_rows(conn, pid)
    flat = [{**a, "fu1_due": (a["fu1"] or {}).get("due"), "fu2_due": (a["fu2"] or {}).get("due")} for a in apps]
    people = [{**c, "app_id": a["id"], "company": a["company"], "role_title": a["role_title"]}
              for a in apps for c in a["contacts"]]
    book = openpyxl.Workbook()
    book.remove(book.active)
    _sheet(book, "Applications", APP_COLUMNS, flat)
    _sheet(book, "Contacts", CONTACT_COLUMNS, people)
    _sheet(book, "Vendors", VENDOR_COLUMNS, vendors.listing(conn, pid))
    book.save(out)
    return out


def export(conn, pid: str, kind: str = "original", out: Path | None = None) -> Path:
    """kind 'original': the person's own layout (a mapped workbook when profiles/<p>/tracker_map.yaml exists, else
    the 14-column tracker); kind 'full': every field (export_full)."""
    from . import tracker, tracker_map
    if kind == "full":
        return export_full(conn, pid, out)
    if kind != "original":
        raise ValueError("kind must be original or full")
    mapping = config.PROFILES / pid / "tracker_map.yaml"
    if mapping.exists():
        return tracker_map.export_mapped(conn, pid, mapping, out)
    return tracker.export_tracker(out, pid, conn=conn)


def summary(path: Path) -> dict:
    """Rows per sheet of a written export (for the CLI and tests)."""
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        return {s.title: s.max_row - 1 for s in book.worksheets}
    finally:
        book.close()
