"""Fictional sample profiles for demo mode and the tests. Nobody here is real: Alex Rivera has four resumes (one per
target role), Jordan Lee has one base resume. The .docx files are generated, never stored in git.

    write(root)   creates root/alex and root/jordan (profile.yaml, resumes/*.docx, baseline.json, known_skills.txt)
"""
import io
import json
from pathlib import Path

import docx
import yaml

from .resume import baseline

ALEX = {
    "name": "Alex Rivera", "headline_default": "Machine Learning Engineer", "email": "alex@example.com",
    "phone": "555-010-0100", "location": "Austin, TX", "linkedin": "linkedin.com/in/alex-rivera-example",
    "github": "github.com/alex-rivera-example", "years": 4,
    "families": ["ai_engineer", "ml_engineer", "data_scientist", "data_engineer"],
    "resume_file_pattern": "{name}_{title}_Resume({company}).pdf",
    "resumes": [
        {"id": "ai_engineer", "label": "AI Engineer", "file": "Alex_Rivera_AI_Engineer_Resume.docx",
         "families": ["ai_engineer", "genai_engineer"]},
        {"id": "ml_engineer", "label": "ML Engineer", "file": "Alex_Rivera_ML_Engineer_Resume.docx",
         "families": ["ml_engineer", "dl_engineer"]},
        {"id": "data_scientist", "label": "Data Scientist", "file": "Alex_Rivera_Data_Scientist_Resume.docx",
         "families": ["data_scientist"]},
        {"id": "data_engineer", "label": "Data Engineer", "file": "Alex_Rivera_Data_Engineer_Resume.docx",
         "families": ["data_engineer"]},
    ],
}
JORDAN = {
    "name": "Jordan Lee", "headline_default": "Data Scientist", "email": "jordan@example.com",
    "location": "Denver, CO", "years": 3, "families": ["data_scientist", "ai_engineer"],
    "resume_file_pattern": "{name}_{title}({company}).pdf",
}
JOBS = [
    ("Northwind Analytics, USA", "Jan 2023 – Present", "Machine Learning Engineer", [
        "Built a retrieval-augmented generation (RAG) assistant with LangChain and PyTorch that answers 4,000 "
        "support questions a week, cutting first-response time by 38%.",
        "Fine-tuned LLM ranking models and served them with FastAPI on Kubernetes (AWS EKS) at 120 ms p95.",
        "Designed offline and online evaluation for model releases, with A/B tests tracked in MLflow.",
        "Moved feature pipelines from cron scripts to Airflow and Spark, saving 11 engineer-hours a week.",
    ]),
    ("Contoso Health, USA", "Jun 2021 – Dec 2022", "Data Scientist", [
        "Trained gradient-boosted readmission models (XGBoost, scikit-learn) used by 14 clinics.",
        "Wrote SQL and Python dashboards in Tableau that replaced a weekly manual Excel report.",
        "Ran causal-impact analysis on a care program and presented results to clinical leadership.",
    ]),
    ("Fabrikam Labs, USA", "Jun 2020 – May 2021", "Data Engineer", [
        "Built ETL jobs in Python and dbt that loaded 2 TB of sensor data into Snowflake each month.",
        "Added data-quality checks with Great Expectations, catching schema drift before it reached reports.",
    ]),
]
SKILLS = [
    ("Languages", "Python, SQL, Bash"),
    ("Machine Learning", "PyTorch, TensorFlow, scikit-learn, XGBoost, LLM, RAG, LangChain, NLP"),
    ("Data", "Spark, Airflow, dbt, Snowflake, Pandas, Tableau"),
    ("Cloud and MLOps", "AWS, Docker, Kubernetes, MLflow, FastAPI, Git"),
]
FOCUS = {
    "ai_engineer": ("AI Engineer", "AI engineer who ships LLM and RAG features to production, from prototype to "
                    "evaluation to serving."),
    "ml_engineer": ("ML Engineer", "Machine learning engineer with 4 years of building, evaluating and serving "
                    "models on AWS and Kubernetes."),
    "data_scientist": ("Data Scientist", "Data scientist who turns messy health and product data into models and "
                       "decisions people act on."),
    "data_engineer": ("Data Engineer", "Data engineer who builds reliable pipelines in Spark, Airflow and dbt for "
                      "analytics and machine learning."),
}
EDUCATION = "M.S. Computer Science | University of Texas at Austin\tMay 2020"


def resume_docx(name: str, contact: str, headline: str, summary: str, skills=SKILLS, jobs=JOBS) -> bytes:
    """A plain resume .docx in the layout baseline.parse reads."""
    doc = docx.Document()
    doc.add_paragraph(name)
    doc.add_paragraph(headline)
    doc.add_paragraph(contact)
    doc.add_paragraph("Summary")
    doc.add_paragraph(summary)
    doc.add_paragraph("Skills")
    for category, items in skills:
        doc.add_paragraph(f"{category}: {items}")
    doc.add_paragraph("Work Experience")
    for company, dates, title, bullets in jobs:
        doc.add_paragraph(f"{company}\t{dates}")
        doc.add_paragraph(title)
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")
    doc.add_paragraph("Education")
    doc.add_paragraph(EDUCATION)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _contact(p: dict) -> str:
    return " | ".join(str(p[k]) for k in ("location", "email", "phone", "linkedin", "github") if p.get(k))


def _write_profile(folder: Path, data: dict) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "profile.yaml").write_text("# Sample profile (fictional). Replace it with your own in Settings.\n"
                                         + yaml.safe_dump(data, sort_keys=False, allow_unicode=True))


def write(root: Path) -> list[str]:
    """Both sample profiles under root (profiles/ of an app home). Returns their ids."""
    root = Path(root)
    alex = root / "alex"
    _write_profile(alex, ALEX)
    (alex / "resumes").mkdir(exist_ok=True)
    for entry in ALEX["resumes"]:
        headline, summary = FOCUS[entry["id"]]
        (alex / "resumes" / entry["file"]).write_bytes(
            resume_docx(ALEX["name"], _contact(ALEX), headline, summary))
    main = alex / "resumes" / ALEX["resumes"][1]["file"]
    parsed = baseline.parse(main)
    (alex / "baseline.json").write_text(json.dumps(parsed, indent=2, ensure_ascii=False) + "\n")
    (alex / "known_skills.txt").write_text("# Skills I really have but my resume does not list.\nGCP\nJAX\n")

    jordan = root / "jordan"
    _write_profile(jordan, JORDAN)
    jobs = [("Initech Insights, USA", "Aug 2022 – Present", "Data Scientist", [
        "Built churn and lifetime-value models in Python and scikit-learn for a subscription product.",
        "Shipped a GPT-based ticket triage prototype with LangChain that routes 60% of tickets automatically.",
        "Owned the experimentation dashboard in Looker and SQL used by three product teams."]),
        ("Hooli Retail, USA", "Jan 2021 – Jul 2022", "Data Analyst", [
            "Automated weekly sales forecasts with Prophet and Airflow.",
            "Cleaned and joined point-of-sale data in BigQuery for the pricing team."])]
    skills = [("Languages", "Python, SQL, R"), ("Machine Learning", "scikit-learn, XGBoost, LLM, LangChain, NLP"),
              ("Data", "BigQuery, Airflow, Looker, Pandas")]
    data = resume_docx(JORDAN["name"], _contact(JORDAN), "Data Scientist",
                       "Data scientist with 3 years of modeling, experimentation and LLM prototypes.", skills, jobs)
    (jordan / "resumes").mkdir(exist_ok=True)
    (jordan / "resumes" / "Jordan_Lee_Data_Scientist.docx").write_bytes(data)
    parsed = baseline.parse(jordan / "resumes" / "Jordan_Lee_Data_Scientist.docx")
    (jordan / "baseline.json").write_text(json.dumps(parsed, indent=2, ensure_ascii=False) + "\n")
    return ["alex", "jordan"]
