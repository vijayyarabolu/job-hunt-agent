"""Import a tracker workbook whose layout differs from the 14-column one (the partner's workbook).

A mapping file (YAML, data not code) names, per sheet, which source column feeds which tracker column, extra
columns kept for the profile's Sheet tab (stored as tagged cells with a 'col' key after the 14 standard ones), a
contact sheet imported as a separate vendors list, how real date cells are read, and explicit date fixes the user
approved. Rows from several sheets are deduplicated by company + role + date before the normal import
(tracker.import_rows), so re-running is idempotent like `tracker import`.

export_mapped writes the profile's tracker back in that workbook's own layout: each row to the sheet it came
from (new rows to the first), original cells unless edited in the app, the vendors sheet from the DB, and the same
column-A "job-hunt-agent #id" note `tracker export` writes. Importing that file again matches every row by its note, and
its dates are not swapped or re-fixed (they were corrected on the way in), so the round trip changes nothing.
"""
import json
import shutil
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
import yaml

from . import config, db
from .text import norm
from . import tracker
from .tracker import COLUMNS, SheetRow, as_text, decode, encode, encode_cell, import_rows, parse_applied

SHEET_ROW_STRIDE = 100_000          # synthetic row numbers: sheet n's row r -> n * stride + r (keeps rows distinct)
DATE_COLUMN = "Date Applied"
MAX_MONTH = 12


@dataclass(frozen=True)
class Mapping:
    sheets: tuple[dict, ...]        # [{name, columns: {source: target}, fill_if_empty: {source: target}}]
    extras: tuple[str, ...]
    vendors: dict | None            # {sheet, columns: {source: field}}
    swap_date_cells: bool
    date_fixes: dict
    template: str | None = None     # the original workbook (relative to the project): export_mapped copies its styles


def load_mapping(path: Path) -> Mapping:
    raw = yaml.safe_load(path.read_text()) or {}
    sheets = tuple(raw.get("sheets") or ())
    if not sheets:
        raise ValueError(f"{path}: no sheets mapped")
    targets = set(COLUMNS) | set(raw.get("extras") or ())
    for sheet in sheets:
        for target in {**sheet.get("columns", {}), **sheet.get("fill_if_empty", {})}.values():
            if target not in targets:
                raise ValueError(f"{path}: {sheet['name']!r} maps to unknown column {target!r}")
    return Mapping(sheets, tuple(raw.get("extras") or ()), raw.get("vendors"), bool(raw.get("swap_date_cells")),
                   {str(k): str(v) for k, v in (raw.get("date_fixes") or {}).items()}, raw.get("template"))


def fix_date(value, mapping: Mapping):
    """A real date cell read with the day and month swapped (sheet formatted dd/mm while typed mm/dd) is swapped
    back when the mapping says so; a typed text date is replaced only by an approved fix."""
    if isinstance(value, datetime):
        if mapping.swap_date_cells and value.day <= MAX_MONTH:
            return value.replace(month=value.day, day=value.month)
        return value
    text = as_text(value)
    return mapping.date_fixes.get(text.strip(), value) if text else value


def _header(sheet) -> dict[str, int]:
    return {str(c.value).strip(): n for n, c in enumerate(sheet[1]) if c.value not in (None, "")}


def _tagged(cell, value) -> dict:
    """The cell tagged for round-trip export; a corrected date keeps the cell's hyperlink but not its old value."""
    tagged = encode_cell(cell)
    return tagged if value is cell.value else {**encode(value), "h": tagged.get("h")}


def _picked(cells, header: dict, spec: dict, mapping: Mapping, exported: bool = False) -> dict[str, tuple]:
    picked: dict[str, tuple] = {}
    for source, target in spec.get("columns", {}).items():
        cell = cells[header[source]]
        fixed = target == DATE_COLUMN and not exported            # our own export holds corrected dates already
        picked[target] = (cell, fix_date(cell.value, mapping) if fixed else cell.value)
    for source, target in spec.get("fill_if_empty", {}).items():
        cell = cells[header[source]]
        if cell.value not in (None, "") and picked.get(target, (None, None))[1] in (None, ""):
            picked[target] = (cell, cell.value)
    return picked


def _sheet_rows(sheet, index: int, spec: dict, mapping: Mapping) -> list[SheetRow]:
    header = _header(sheet)
    missing = [s for s in {**spec.get("columns", {}), **spec.get("fill_if_empty", {})} if s not in header]
    if missing:
        raise ValueError(f"sheet {sheet.title!r} has no column {missing}")
    rows = []
    for cells in sheet.iter_rows(min_row=2):
        note = tracker._note_match(cells[0])
        picked = _picked(cells, header, spec, mapping, exported=note is not None)
        values = [picked[c][1] if c in picked else None for c in COLUMNS]
        if all(v in (None, "") for v in values[:2]):
            continue                    # no company and no role: an empty or date-only (pre-filled) row
        tagged = [_tagged(*picked[c]) if c in picked else encode(None) for c in COLUMNS]
        tagged += [{**_tagged(*picked[e]), "col": e} for e in mapping.extras if e in picked]
        mapped = spec.get("columns", {})                 # her other columns ride along unchanged
        tagged += [{**encode_cell(cells[i]), "col": head} for head, i in header.items()
                   if head not in mapped and i < len(cells) and cells[i].value not in (None, "")]
        rows.append(SheetRow(index * SHEET_ROW_STRIDE + cells[0].row, values, tagged,
                             int(note.group(1)) if note else None, note.group(2) if note else None))
    return rows


def identity(row: SheetRow) -> tuple[str, str, str]:
    company, role, applied = row.values[:3]
    return norm(as_text(company)), norm(as_text(role)), parse_applied(applied) or norm(as_text(applied))


def merged_rows(workbook, mapping: Mapping) -> tuple[list[SheetRow], dict]:
    """All mapped rows; a row whose company + role + date already came from an EARLIER sheet is dropped (the same
    application tracked in two sheets). Repeats inside one sheet stay, as `tracker import` keeps them."""
    out, seen = [], {}
    stats = {"per_sheet": {}, "cross_sheet_duplicates": 0, "within_sheet_repeats": 0}
    for index, spec in enumerate(mapping.sheets, 1):
        rows = _sheet_rows(workbook[spec["name"]], index, spec, mapping)
        stats["per_sheet"][spec["name"]] = len(rows)
        for row in rows:
            key = identity(row)
            if seen.get(key, index) != index:
                stats["cross_sheet_duplicates"] += 1
                continue
            stats["within_sheet_repeats"] += key in seen
            seen[key] = index
            out.append(row)
    return out, stats


def vendors(workbook, mapping: Mapping) -> list[dict]:
    spec = mapping.vendors
    if not spec:
        return []
    sheet = workbook[spec["sheet"]]
    header = _header(sheet)
    missing = [s for s in spec["columns"] if s not in header]
    if missing:
        raise ValueError(f"sheet {sheet.title!r} has no column {missing}")
    out = []
    for cells in sheet.iter_rows(min_row=2):
        entry = {field: as_text(cells[header[src]].value) for src, field in spec["columns"].items()}
        if any(entry.values()):
            out.append(entry)
    return out


def import_mapped(xlsx: Path, profile: str, mapping_path: Path, conn, dry: bool = False) -> dict:
    """Import the workbook into the profile's applications; vendors go to profiles/<p>/vendors.json (gitignored)."""
    mapping = load_mapping(mapping_path)
    workbook = openpyxl.load_workbook(xlsx)
    rows, stats = merged_rows(workbook, mapping)
    found = vendors(workbook, mapping)
    unparsed = sorted({as_text(r.values[2]) or "" for r in rows if not parse_applied(r.values[2])})
    result = {"profile": profile, "rows": len(rows), **stats, "vendors": len(found), "unparsed_dates": unparsed}
    if dry:
        return result
    target = config.PROFILES / profile / "vendors.json"
    target.write_text(json.dumps(found, indent=2, ensure_ascii=False) + "\n")
    target.chmod(0o600)
    imported = import_rows(rows, profile, conn)
    from . import tracker_layout, vendors as vendor_list         # the DB list (kept when edited in the app)
    result["vendors_replaced"] = vendor_list.from_import(conn, profile, found)
    index = {spec["name"]: n for n, spec in enumerate(mapping.sheets, 1)}
    with db.tx(conn):                                            # app fields from our own export
        result["extras_restored"] = tracker_layout.apply_extras(
            conn, profile, tracker_layout.read_extras(workbook),
            lambda tab, row: index.get(tab, 1) * SHEET_ROW_STRIDE + row)
    return {**result, **{k: v for k, v in imported.items() if k not in ("profile", "rows")}}


# ---------- export in the workbook's own layout ----------

EXTRA_FIELD = {"Address": "address", "Team Name": "team_name"}


def _extra_cell(app: dict, header: str):
    """Address / Team Name: the app's value when edited there, else the imported cell (type and link kept)."""
    field = EXTRA_FIELD.get(header)
    if field and field in json.loads(app["edited"] or "[]"):
        return app[field], None
    raw = json.loads(app["raw_json"]) if app["raw_json"] else []
    tagged = next((c for c in raw if isinstance(c, dict) and c.get("col") == header), None)
    return (decode(tagged), tagged.get("h")) if tagged else (app.get(field) if field else None, None)


def _sheet_index(app: dict, n_sheets: int) -> int:
    index = (app["xlsx_row"] or 0) // SHEET_ROW_STRIDE
    return index if 1 <= index <= n_sheets else 1


def _clear(sheet) -> None:
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.hyperlink = None
            cell.value = None
        if tracker._note_id(row[0]) is not None:
            row[0].comment = None


def _book(mapping: Mapping, template: Path | None, out: Path):
    if template and template.exists() and template.resolve() != out.resolve():
        shutil.copyfile(template, out)
        return openpyxl.load_workbook(out)
    book = openpyxl.Workbook()
    book.remove(book.active)
    for spec in mapping.sheets:
        book.create_sheet(spec["name"]).append([*spec.get("columns", {}), *spec.get("fill_if_empty", {})])
    if mapping.vendors:
        book.create_sheet(mapping.vendors["sheet"]).append(list(mapping.vendors["columns"]))
    return book


def _guarded(cell):
    """Text that starts like a formula (= + - @) stays text in Excel."""
    if isinstance(cell.value, str) and cell.value.startswith(tracker._FORMULA_CHARS):
        cell.data_type = "s"
        cell.quotePrefix = True
    return cell


def _write_vendors(book, mapping: Mapping, entries: list[dict]) -> None:
    spec = mapping.vendors
    sheet = book[spec["sheet"]]
    header = _header(sheet)
    _clear(sheet)
    for n, entry in enumerate(entries, 2):
        for source, field in spec["columns"].items():
            _guarded(sheet.cell(row=n, column=header[source] + 1, value=entry.get(field)))


def export_mapped(conn, profile: str, mapping_path: Path, out: Path | None = None,
                  template: Path | None = None) -> Path:
    """The profile's tracker in its own workbook layout (see the module docstring). Returns the file written."""
    from . import vendors as vendor_list
    mapping = load_mapping(mapping_path)
    out = Path(out or config.DATA / "export" / f"Application_Tracker_{profile}_{date.today():%Y%m%d}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    template = template or (config.ROOT / mapping.template if mapping.template else None)
    book = _book(mapping, template, out)
    sheets = [book[spec["name"]] for spec in mapping.sheets]
    headers = [_header(sheet) for sheet in sheets]
    for sheet in sheets:
        _clear(sheet)
    from . import trackapi, tracker_layout
    layout = tracker_layout.of_mapping(profile, mapping_path)
    apps = [dict(r) for r in conn.execute("SELECT * FROM applications WHERE profile = ? "
                                          "ORDER BY coalesce(xlsx_row, 1e12), id", (profile,))]
    next_row, extra_rows, today = [2] * len(sheets), [], date.today()
    for app in apps:
        k = _sheet_index(app, len(sheets)) - 1
        row = next_row[k]
        next_row[k] += 1
        for column in layout.app_tabs[k].columns:                # every column of her tab, in her order
            if column.header not in headers[k]:
                continue
            value, link, _ = tracker_layout.cell(app, column, layout)
            cell = _guarded(sheets[k].cell(row=row, column=headers[k][column.header] + 1, value=value))
            tracker._link(cell, link)
        tracker._note(sheets[k].cell(row=row, column=1), app["id"], app["status"])
        dues = {"fu1": trackapi._due(app, 1, today), "fu2": trackapi._due(app, 2, today)}
        extra_rows.append((sheets[k].title, row, {**app, **tracker_layout.extras(app, dues)}))
    tracker_layout.write_extras(book, extra_rows)
    if mapping.vendors:
        _write_vendors(book, mapping, vendor_list.listing(conn, profile))
    book.save(out)
    return out
