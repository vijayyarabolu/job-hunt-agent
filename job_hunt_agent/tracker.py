"""Application tracker: lossless xlsx import/export (GOAL §6), pipeline status, follow-ups."""
import difflib
import functools
import json
import re
import shutil
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.worksheet.hyperlink import Hyperlink

from . import config, db, history
from .applied import company_key, same_company
from .google import state as gstate             # status/field edits mark the Sheet dirty (no network)
from .text import norm, sha1

COLUMNS = ["Company Name", "Role Title", "Date Applied", "JD", "Resume", "Target Name", "Target Role",
           "LinkedIn Request", "Email Address", "Email Sent", "Gemini chat number", "Follow-Up 1(3days)",
           "Follow-Up 2(5-7days)", "Status(A/R)"]
FIELDS = ["company", "role_title", "date_applied", "jd", "resume", "target_name", "target_role",
          "linkedin_request", "email_address", "email_sent", "gemini_chat", "follow_up_1", "follow_up_2", "status_ar"]
STATUSES = ("saved", "tailored", "applied", "screen", "interview", "offer", "rejected", "ghosted")
# Status(A/R) shows the status as people type it (Jordan's "Rejected"); ghosted is "Cold" there
STATUS_WORDS = {"saved": "Saved", "tailored": "Tailored", "applied": "Applied", "screen": "Screen",
                "interview": "Interview", "offer": "Offer", "rejected": "Rejected", "ghosted": "Cold"}
SUPPRESS_FOLLOWUPS_AFTER_DAYS = 21
FOLLOWUP_DAYS = (3, 6)
_FORMULA_CHARS = ("=", "+", "-", "@")
IDENTITY = ("company", "role_title", "date_applied")
NOTE_AUTHOR = "job-hunt-agent"
_NOTE = re.compile(r"\s*job-hunt-agent #(\d+)(?:\s*·\s*([a-z]+))?\s*$")     # "job-hunt-agent #375 · tailored"


def default_import_path() -> Path:
    return config.DATA / "import" / "Application_Tracker.xlsx"


def encode(value) -> dict:
    """Tag a cell value with its type so export can write back the identical cell."""
    if value is None:
        return {"t": "none"}
    if isinstance(value, bool):
        return {"t": "bool", "v": value}
    if isinstance(value, datetime):
        return {"t": "dt", "v": value.isoformat()}
    if isinstance(value, date):
        return {"t": "date", "v": value.isoformat()}
    if isinstance(value, (int, float)):
        return {"t": "num", "v": value}
    return {"t": "str", "v": str(value)}


def encode_cell(cell) -> dict:
    """Like encode(), but a real formula cell is tagged 'f' so it round-trips as a formula, not as text, and the
    cell's hyperlink travels with it ('h'), so a sorted sheet never leaves links on other rows (verifier M9)."""
    tagged = {"t": "f", "v": str(cell.value)} if cell.data_type == "f" else encode(cell.value)
    link = cell.hyperlink
    return {**tagged, "h": {"target": link.target, "location": link.location} if link else None}


def decode(tagged: dict):
    kind, value = tagged["t"], tagged.get("v")
    if kind == "dt":
        return datetime.fromisoformat(value)
    if kind == "date":
        return date.fromisoformat(value)
    return None if kind == "none" else value


def as_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def parse_applied(value) -> str | None:
    """Date Applied (datetime or 'MM/DD/YYYY'-ish string) -> ISO date, or None if unparseable."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).strftime("%Y-%m-%d")
        except (TypeError, ValueError):
            continue
    return None


def import_key(profile: str, company, role, applied) -> str:
    return sha1(profile, norm(as_text(company)), norm(as_text(role)), parse_applied(applied) or f"raw:{as_text(applied)}")


def _profile(profile: str | None) -> str:
    profile = profile or config.default_profile()
    if not profile:
        raise ValueError("no profile yet: create profiles/<id>/profile.yaml (scripts/setup.sh does)")
    return profile


def ensure_profile(conn, profile_id: str) -> None:
    name = profile_id
    try:
        name = config.profile(profile_id)["name"]
    except config.ConfigError:
        pass
    with conn:
        conn.execute("INSERT OR IGNORE INTO profiles(id, name) VALUES (?,?)", (profile_id, name))


@dataclass(frozen=True)
class SheetRow:
    number: int                 # 1-based sheet row
    values: list
    tagged: list
    note_id: int | None         # application id from the column-A note export writes (moves with Excel sorts)
    note_status: str | None = None  # the app's status when THIS file was exported (same note)
    key: str = ""


def _note_match(cell):
    return _NOTE.match((cell.comment.text if cell.comment else "") or "")


def _note_id(cell) -> int | None:
    match = _note_match(cell)
    return int(match.group(1)) if match else None


def _read_rows(path: Path) -> tuple[list[SheetRow], str, dict]:
    """(rows, tab name, extras sheet values) of a tracker workbook: the 14 columns in any order, plus columns the app
    added to the real sheet (Address), which are kept as tagged extra cells; raw_json stays positional."""
    from . import tracker_layout
    book = openpyxl.load_workbook(path)
    sheet = book.worksheets[0]
    header = [str(c.value).strip() if c.value is not None else "" for c in sheet[1]]
    if not set(COLUMNS) <= set(header):
        raise ValueError(f"unexpected tracker header {header}; expected the columns {COLUMNS}")
    at = [header.index(c) for c in COLUMNS]
    added = [(h, i) for i, h in enumerate(header) if h in EXTRA_COLUMNS]
    rows = []
    for row in sheet.iter_rows(min_row=2, max_col=len(header)):
        cells = [row[i] for i in at]
        extra = [(h, row[i]) for h, i in added]
        if any(c.value not in (None, "") for c in [*cells, *(c for _, c in extra)]):
            note = _note_match(row[0])
            tagged = [encode_cell(c) for c in cells] + [{**encode_cell(c), "col": h} for h, c in extra]
            rows.append(SheetRow(row[0].row, [c.value for c in cells], tagged,
                                 int(note.group(1)) if note else None, note.group(2) if note else None))
    return rows, sheet.title, tracker_layout.read_extras(book)


def _with_keys(profile: str, rows: list[SheetRow]) -> list[SheetRow]:
    """import_key per row; repeated identical rows get #2, #3 in sheet order."""
    seen: dict[str, int] = {}
    out = []
    for row in rows:
        key = import_key(profile, *row.values[:3])
        seen[key] = seen.get(key, 0) + 1
        out.append(replace(row, key=key if seen[key] == 1 else f"{key}#{seen[key]}"))
    return out


def _identity_score(app, values: list) -> int:
    return sum(norm(app[f]) == norm(as_text(v)) for f, v in zip(IDENTITY, values[:3]))


def _base(key: str | None) -> str:
    return (key or "").split("#", 1)[0]


def _match_keys(rows: list[SheetRow], apps: dict[int, dict]) -> dict[int, int]:
    """Rows and apps with the same identity (company, role, date) pair up: same sheet row first, then the exact
    key, then in order. Identical repeated rows therefore never trade places when another one is edited."""
    pairs: dict[int, int] = {}
    groups: dict[str, list[int]] = {}
    for i, a in apps.items():
        if a["import_key"]:
            groups.setdefault(_base(a["import_key"]), []).append(i)
    for rule in (lambda r, a: a["xlsx_row"] == r.number, lambda r, a: a["import_key"] == r.key, lambda r, a: True):
        for row in rows:
            free = [i for i in groups.get(_base(row.key), ()) if i not in pairs.values()]
            match = next((i for i in free if rule(row, apps[i])), None)
            if row.number not in pairs and match is not None:
                pairs[row.number] = match
    return pairs


def _one_row_per_note(rows: list[SheetRow], apps: dict[int, dict]) -> tuple[list[SheetRow], list[int]]:
    """Excel copies a note with its row. When rows share an application's note, one row keeps it, ranked by: company,
    role and date matching the application; the closest company name; most cells still as that file was exported;
    its previous sheet row; topmost. The other rows are matched like note-free rows (usually new applications), so a
    copy never silently takes over its source. Rows that tied with the keeper on identity are returned as conflicts
    so the user sees the choice (verifier M9 rounds 3-5)."""
    groups: dict[int, list[SheetRow]] = {}
    for row in rows:
        if row.note_id in apps:
            groups.setdefault(row.note_id, []).append(row)
    keep: dict[int, int] = {}
    tied: list[int] = []
    for app_id, group in groups.items():
        app = apps[app_id]
        exported = _exported_text(app, group[0].note_status)
        ranked = sorted(group, reverse=True, key=lambda r: (
            _identity_score(app, r.values),
            _company_similarity(app["company"], as_text(r.values[0])),
            sum(norm(a) == norm(as_text(b)) for a, b in zip(exported, r.values)),
            r.number == app["xlsx_row"], -r.number))
        keep[app_id] = ranked[0].number
        best = _identity_score(app, ranked[0].values)
        tied += [r.number for r in ranked[1:] if _identity_score(app, r.values) == best]
    kept = [row if row.note_id not in keep or keep[row.note_id] == row.number
            else replace(row, note_id=None, note_status=None) for row in rows]
    return kept, sorted(tied)


def _exported_text(app: dict, then: str | None = None) -> list[str | None]:
    """The 14 cells as export wrote them for this application, as text; the status cell as of the file's own export
    (`then`, from its note) when known, since the app may have moved on since (verifier M9 round 5)."""
    values = [_cell_value(app, idx)[0] for idx in range(len(COLUMNS))]
    status_idx = FIELDS.index("status_ar")
    values[status_idx] = _status_cell(values[status_idx], {**app, "status": then} if then in STATUSES else app)
    return [as_text(v) for v in values]


def _company_similarity(a: str | None, b: str | None) -> float:
    """1.0 for the same company (spelling aside), else how close the names are ("Deloite" vs "Deloitte LLP")."""
    x, y = company_key(a), company_key(b)
    if same_company(x, y):
        return 1.0
    return difflib.SequenceMatcher(None, x.replace(" ", ""), y.replace(" ", "")).ratio() if x and y else 0.0


def _match(rows: list[SheetRow], apps: dict[int, dict]) -> tuple[dict[int, int], list[int]]:
    """sheet row number -> application id. By the export note first, then import_key, then (an identity field
    edited without a note) the one unclaimed application with 2 of 3 identity fields equal; a tie is broken only by
    the same sheet row, otherwise the row is a conflict and imported as new (release review C2)."""
    claimed: dict[int, int] = {}
    for row in rows:
        if row.note_id in apps and row.note_id not in claimed.values():
            claimed[row.number] = row.note_id
    claimed.update(_match_keys([r for r in rows if r.number not in claimed],
                               {i: a for i, a in apps.items() if i not in claimed.values()}))
    conflicts = []
    for row in rows:
        if row.number in claimed:
            continue
        taken = set(claimed.values())
        cands = [i for i, a in apps.items() if i not in taken and _identity_score(a, row.values) >= 2]
        if len(cands) > 1:
            cands = [i for i in cands if apps[i]["xlsx_row"] == row.number] or cands
        if len(cands) == 1:
            claimed[row.number] = cands[0]
        elif cands:
            conflicts.append(row.number)
    return claimed, conflicts


def _kept(row: SheetRow, app: dict, text: dict, keep: set[str]) -> list[dict]:
    """Sheet values that differ from a field the app edited (the app's value wins; the user is told)."""
    return [{"row": row.number, "field": f, "sheet": text[f], "app": app[f]}
            for f in FIELDS if f in keep and norm(text[f]) != norm(app[f])]


def _update(conn, row: SheetRow, app: dict, text: dict, stale: bool) -> list[dict]:
    keep = set(json.loads(app["edited"] or "[]"))
    sets = {f: v for f, v in text.items() if f not in keep}
    sets["followups_suppressed"] = int(stale)
    sets.update({k: v for k, v in _extras(row.tagged).items() if k not in keep})
    applied_on = parse_applied(row.values[2])
    if "date_applied" not in keep and applied_on:                   # a date fixed in the sheet moves follow-ups
        sets["applied_on"] = applied_on
    if not conn.execute("SELECT 1 FROM applications WHERE import_key=? AND id<>?", (row.key, app["id"])).fetchone():
        sets["import_key"] = row.key
    assignments = ", ".join(f"{f}=?" for f in sets)
    conn.execute(f"UPDATE applications SET {assignments}, raw_json=?, xlsx_row=? WHERE id=?",
                 [*sets.values(), json.dumps(row.tagged), row.number, app["id"]])
    then = row.note_status if row.note_id == app["id"] and row.note_status in STATUSES else None
    return _kept(row, app, text, keep) + _status_from_sheet(conn, app, row.number, text["status_ar"], then)


EXTRA_COLUMNS = {"Address": "address", "Team Name": "team_name"}     # mapped-import extras with their own field


def _extras(tagged: list[dict]) -> dict[str, str | None]:
    """Jordan's extra cells (tagged with 'col' by tracker_map) as their GOAL3 fields."""
    found = {c["col"]: as_text(decode(c)) for c in tagged if isinstance(c, dict) and c.get("col") in EXTRA_COLUMNS}
    return {field: found.get(col) for col, field in EXTRA_COLUMNS.items() if col in found}


def _insert(conn, profile: str, row: SheetRow, text: dict, stale: bool) -> None:
    applied_on = parse_applied(row.values[2])
    key = row.key
    if conn.execute("SELECT 1 FROM applications WHERE import_key=?", (key,)).fetchone():
        key = None                                                  # a conflict row: its key belongs to another app
    status = _sheet_status(text["status_ar"]) or default_status(applied_on)
    extras = _extras(row.tagged)
    cur = conn.execute(
        f"INSERT INTO applications(profile, import_key, {', '.join(FIELDS)}, applied_on, status, raw_json, "
        f"xlsx_row, sheet_status, followups_suppressed, updated_at, source{''.join(', ' + k for k in extras)}) "
        f"VALUES (?,?,{','.join('?' * len(FIELDS))},?,?,?,?,?,?,?,'import'{',?' * len(extras)})",
        [profile, key, *text.values(), applied_on, status, json.dumps(row.tagged), row.number,
         _sheet_status(text["status_ar"]) or "", int(stale), db.iso(db.utcnow()), *extras.values()])
    conn.execute("INSERT INTO application_events(app_id, at, from_status, to_status, source) "
                 "VALUES (?,?,?,?,?)", (cur.lastrowid, db.iso(db.utcnow()), None, status, "import"))


def _stale(applied_on: str | None, today: date) -> bool:
    return applied_on is not None and not (
        today - timedelta(days=SUPPRESS_FOLLOWUPS_AFTER_DAYS) <= date.fromisoformat(applied_on) <= today)


def import_tracker(path: Path | None = None, profile: str | None = None, conn=None) -> dict:
    """Idempotent: re-import updates matched rows (see _match) and never duplicates; app edits are kept and any
    sheet value they override is listed in kept_app_values. The DB is the source of truth; the xlsx is its view.
    Follow-ups are suppressed for rows applied > 21 days ago or dated in the future (likely day/month typos)."""
    from . import tracker_layout
    conn = conn or db.connect()
    profile = _profile(profile)
    rows, _, extras = _read_rows(Path(path or default_import_path()))
    out = import_rows(rows, profile, conn)
    with db.tx(conn):
        out["extras_restored"] = tracker_layout.apply_extras(conn, profile, extras, lambda tab, row: row)
    return out


def import_rows(rows: list[SheetRow], profile: str, conn) -> dict:
    """The import itself, for rows from the user's own tracker or from a mapped workbook (tracker_map.py)."""
    ensure_profile(conn, profile)
    rows = _with_keys(profile, rows)
    apps = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM applications WHERE profile=?", (profile,))}
    rows, tied = _one_row_per_note(rows, apps)
    claimed, conflicts = _match(rows, apps)
    conflicts = sorted(set(conflicts) | set(tied))
    inserted = updated = 0
    kept: list[dict] = []
    before = conn.execute("SELECT coalesce(max(id), 0) FROM applications").fetchone()[0]
    with db.tx(conn):
        for row in rows:
            text = dict(zip(FIELDS, (as_text(v) for v in row.values)))
            stale = _stale(parse_applied(row.values[2]), date.today())
            if row.number in claimed:
                kept += _update(conn, row, apps[claimed[row.number]], text, stale)
                updated += 1
            else:
                _insert(conn, profile, row, text, stale)
                inserted += 1
    from . import contacts                                       # late: contacts imports this module
    touched = [*(claimed[r.number] for r in rows if r.number in claimed),
               *(r["id"] for r in conn.execute("SELECT id FROM applications WHERE id > ? AND profile = ?",
                                               (before, profile)))]
    for i in range(0, len(touched), contacts.SQL_BATCH):         # packed contact cells may have changed
        contacts.refresh_legacy(conn, touched[i:i + contacts.SQL_BATCH])
    return {"profile": profile, "rows": len(rows), "inserted": inserted, "updated": updated,
            "conflicts": conflicts, "kept_app_values": kept}


def default_status(applied_on: str | None) -> str:
    return "applied" if applied_on else "saved"


def _sheet_status(cell: str | None) -> str | None:
    value = (cell or "").strip().lower()
    value = "ghosted" if value == "cold" else value
    return value if value in STATUSES else None


def _status_from_sheet(conn, app: dict, row_number: int, cell: str | None, then: str | None = None) -> list[dict]:
    """A pipeline status the user typed into Status (A/R) becomes the app's status, but an old file never moves it
    back silently (verifier M9): with a note, what the app was when that file was exported (`then`) decides what is
    an edit and whether the app changed since; without one (the user's own tracker), a status the app already moved
    past is reported, not applied. Conflicts go to kept_app_values; the app wins. sheet_status only records what
    the sheet last showed."""
    wanted = _sheet_status(cell)
    conn.execute("UPDATE applications SET sheet_status=? WHERE id=?", (wanted or "", app["id"]))
    if wanted is None or wanted == app["status"]:
        return []
    conflict = [{"row": row_number, "field": "status_ar", "sheet": wanted, "app": app["status"]}]
    if then is not None:                                        # the file says what the app was when it was made
        if wanted == then:                                      # unchanged since: not an edit
            return []
        if app["status"] != then:                               # the app changed since that file: app wins
            return conflict
    elif wanted in _left_statuses(conn, app["id"]):             # no note: a status the app already moved past
        return conflict                                         # (an old copy of the sheet), never silently back
    now = db.iso(db.utcnow())
    conn.execute("UPDATE applications SET status=?, updated_at=? WHERE id=?", (wanted, now, app["id"]))
    conn.execute("INSERT INTO application_events(app_id, at, from_status, to_status, source) VALUES (?,?,?,?,?)",
                 (app["id"], now, app["status"], wanted, "import"))
    return []


def _left_statuses(conn, app_id: int) -> set[str]:
    """Statuses this application had before its current one (from its event history)."""
    rows = conn.execute("SELECT from_status, to_status FROM application_events WHERE app_id=? ORDER BY id",
                        (app_id,)).fetchall()
    had = {s for r in rows for s in (r["from_status"], r["to_status"]) if s}
    current = rows[-1]["to_status"] if rows else None
    return had - {current}


PARSED_KEPT = 8192        # GOAL30-app: rows' stored JSON, parsed once (the tracker list parsed it once per cell)


@functools.lru_cache(maxsize=PARSED_KEPT)
def stored_cells(raw_json: str) -> list:
    """A row's raw_json, parsed once per text. Shared: read it, never change it."""
    return json.loads(raw_json)


@functools.lru_cache(maxsize=PARSED_KEPT)
def edited_fields(edited: str | None) -> frozenset:
    return frozenset(json.loads(edited or "[]"))


def _cell_value(app: dict, idx: int):
    """Original typed cell unless the app edited that field. Returns (value, write_as_text)."""
    field = FIELDS[idx]
    raw = stored_cells(app["raw_json"]) if app["raw_json"] else None
    if raw and field not in edited_fields(app["edited"]):
        tagged = raw[idx]
        if tagged["t"] == "f":
            return tagged["v"], False
        value = decode(tagged)
    else:
        value = app[field]
    return value, isinstance(value, str) and value.startswith(_FORMULA_CHARS)


ADDED_COLUMN_WIDTH = 30


def _template_book(layout, template: Path | None, out: Path):
    """A copy of the person's real workbook (formatting, data validation, table) with the columns the app adds
    inserted where the layout says; a plain workbook with the layout's header when there is no template."""
    from . import tracker_layout
    tab = layout.app_tabs[0]
    if template and template.exists() and template.resolve() != out.resolve():
        shutil.copyfile(template, out)
        book = openpyxl.load_workbook(out)
        sheet = book.worksheets[0]
        for at, column in enumerate(tab.columns, 1):
            if sheet.cell(row=1, column=at).value != column.header:
                tracker_layout.insert_column(sheet, at, column.header, ADDED_COLUMN_WIDTH)
        return book
    book = openpyxl.Workbook()
    book.active.title = tab.name
    book.active.append(tab.headers)
    return book


def export_tracker(out: Path | None = None, profile: str | None = None, template: Path | None = None,
                   conn=None) -> Path:
    """Write the tracker in the person's own layout (their real sheet's columns in order, plus the columns
    the app adds such as Address), reusing their real workbook as the template so formatting, data validation and
    the table survive. The app's own fields go to a separate 'job-hunt-agent app fields' sheet."""
    from . import trackapi, tracker_layout
    conn = conn or db.connect()
    profile = _profile(profile)
    layout = tracker_layout.of_profile(profile, template)
    tab = layout.app_tabs[0]
    template = layout.template
    out = Path(out or config.DATA / "export" / f"Application_Tracker_{profile}_{date.today():%Y%m%d}.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    book = _template_book(layout, template, out)
    sheet = book.worksheets[0]
    template_links = {}
    for row in sheet.iter_rows(min_row=2, max_col=len(tab.columns)):
        for cell in row:
            if cell.hyperlink:
                template_links[(cell.row, cell.column)] = {"target": cell.hyperlink.target,
                                                           "location": cell.hyperlink.location}
            cell.hyperlink = None
            cell.value = None
        if _note_id(row[0]) is not None:                            # our id note; a user's own note stays
            row[0].comment = None
    by_field = {c.field: i for i, c in enumerate(tab.columns) if c.field}
    links = {(r, FIELDS.index(c.field) + 1): v for (r, col), v in template_links.items()
             for c in [tab.columns[col - 1]] if c.field in FIELDS}            # template link by row + 14-col index
    apps = conn.execute("SELECT * FROM applications WHERE profile=? ORDER BY coalesce(xlsx_row, 1e9), id",
                        (profile,)).fetchall()
    seen, stamps, extra_rows, today = {}, [], [], date.today()
    for offset, row_ in enumerate(apps):
        app, number = dict(row_), 2 + offset
        values = []
        for col, column in enumerate(tab.columns, 1):
            value, link, as_text = tracker_layout.cell(app, column, layout, links)
            values.append(value)
            cell = sheet.cell(row=number, column=col, value=value)
            if as_text:
                cell.data_type = "s"
                cell.quotePrefix = True
            _link(cell, link)
        _note(sheet.cell(row=number, column=1), app["id"], app["status"])
        ident = [values[by_field[f]] for f in IDENTITY]
        key = import_key(profile, *ident)                                  # what import will compute
        seen[key] = seen.get(key, 0) + 1
        stamps.append((key if seen[key] == 1 else f"{key}#{seen[key]}", number, app["id"], app["import_key"],
                       _sheet_status(values[by_field["status_ar"]]) or ""))
        dues = {"fu1": trackapi._due(app, 1, today), "fu2": trackapi._due(app, 2, today)}
        extra_rows.append((tab.name, number, {**app, **tracker_layout.extras(app, dues)}))
    tracker_layout.write_extras(book, extra_rows)
    book.save(out)
    _stamp(conn, stamps)
    return out


def _cell_link(app: dict, idx: int, template_links: dict) -> dict | None:
    """The hyperlink the row's cell had when imported; none for a field edited in the app. Rows imported before
    links were recorded keep the template's link at their own sheet row."""
    if FIELDS[idx] in edited_fields(app["edited"]) or not app["raw_json"]:
        return None
    tagged = stored_cells(app["raw_json"])[idx]
    if "h" in tagged:
        return tagged["h"]
    return template_links.get((app["xlsx_row"], idx + 1))


def _link(cell, link: dict | None) -> None:
    if not link:
        return
    value = cell.value
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=link.get("target"), location=link.get("location"))
    cell.value = value                                           # openpyxl fills an empty cell with the URL


def _note(cell, app_id: int, status: str) -> None:
    """The row's application id and current status as a column-A note: it moves with the row when the sheet is
    sorted or filtered, so import keeps every row's identity without a 15th column and knows what this file showed.
    A note the user wrote there is left alone."""
    if cell.comment is None:
        cell.comment = Comment(f"job-hunt-agent #{app_id} · {status}", NOTE_AUTHOR, height=20, width=140)


def _status_cell(value, app) -> str | None:
    """Status (A/R) shows the pipeline status beyond the default, never over text the user typed."""
    if value not in (None, "") and _sheet_status(str(value)) in (None, app["status"]):
        return value                                             # user text, or the same status as they typed it
    if app["status"] == default_status(app["applied_on"]) and value in (None, ""):
        return value
    return STATUS_WORDS.get(app["status"], app["status"])


def _stamp(conn, stamps: list) -> None:
    """Give rows created in the app the import_key and sheet row their exported cells will produce, so the next
    import updates them instead of adding a copy (release review C2), and remember the Status(A/R) each row shows."""
    with db.tx(conn):
        for key, row, app_id, current, shown in stamps:
            taken = conn.execute("SELECT 1 FROM applications WHERE import_key=? AND id<>?", (key, app_id)).fetchone()
            if current is None and taken is None:
                conn.execute("UPDATE applications SET import_key=? WHERE id=?", (key, app_id))
            conn.execute("UPDATE applications SET xlsx_row=?, sheet_status=? WHERE id=?", (row, shown, app_id))


def set_status(conn, app_id: int, status: str, source: str = "cli", snapshot: dict | None = None) -> None:
    if status not in STATUSES:
        raise ValueError(f"status must be one of {STATUSES}")
    row = conn.execute("SELECT status FROM applications WHERE id=?", (app_id,)).fetchone()
    if row is None:
        raise ValueError(f"no application {app_id}")
    now = db.iso(db.utcnow())
    current = dict(conn.execute("SELECT applied_on, edited, date_applied FROM applications WHERE id=?",
                                (app_id,)).fetchone())
    with db.tx(conn):
        updates = {"status": status, "updated_at": now}
        if status == "applied" and not current["applied_on"]:
            updates["applied_on"] = date.today().isoformat()
            updates["date_applied"] = date.today().strftime("%m/%d/%Y")
            updates["edited"] = json.dumps(sorted(set(json.loads(current["edited"] or "[]")) | {"date_applied"}))
        if source != "sheet":
            updates["field_at"] = _stamped(conn, app_id, "status", now)
        conn.execute(f"UPDATE applications SET {', '.join(f'{k}=?' for k in updates)} WHERE id=?",
                     [*updates.values(), app_id])
        conn.execute("INSERT INTO application_events(app_id, at, from_status, to_status, source, snapshot) "
                     "VALUES (?,?,?,?,?,?)", (app_id, now, row["status"], status, source,
                                              json.dumps(snapshot) if snapshot else None))
        history.record(conn, app_id, "status", row["status"], status, "sheet" if source == "sheet" else "app")
        if "date_applied" in updates:
            history.record(conn, app_id, "date_applied", current["date_applied"], updates["date_applied"],
                           "sheet" if source == "sheet" else "app")
        gstate.mark_dirty(conn)


def _stamped(conn, app_id: int, field: str, now: str) -> str:
    """applications.field_at with `field` stamped now: the app-side time the Sheet's last-writer-wins compares
. Edits pulled from the Sheet are not stamped (source='sheet')."""
    row = conn.execute("SELECT field_at FROM applications WHERE id=?", (app_id,)).fetchone()
    return json.dumps({**json.loads(row["field_at"] or "{}"), field: now})


def track_job(conn, profile: str, uid: str, source: str = "ui") -> int:
    """The application row for a fetched job (created as 'saved' on first use, prefilled with the posting's facts:
    company, role, JD link, apply link, location, workplace, salary, employment type). Returns its id."""
    found = conn.execute("SELECT id FROM applications WHERE profile=? AND job_uid=?", (profile, uid)).fetchone()
    if found:
        return found["id"]
    from .trackapi import job_facts                              # late: trackapi imports this module
    facts = job_facts(conn, uid)
    if not facts:
        raise ValueError(f"no job {uid}")
    ensure_profile(conn, profile)
    now = db.iso(db.utcnow())
    cols = {**facts, "jd": facts["job_url"], "source": "feed"}
    with db.tx(conn):
        # GOAL33 §3: looked for again inside the write lock: Autofill's Mark applied and the page's "Did you apply?"
        # can save the same job at the same moment
        found = conn.execute("SELECT id FROM applications WHERE profile=? AND job_uid=?", (profile, uid)).fetchone()
        if found:
            return found["id"]
        cur = conn.execute(
            f"INSERT INTO applications(profile, job_uid, {', '.join(cols)}, status, updated_at) "
            f"VALUES (?,?{',?' * len(cols)}, 'saved', ?)", (profile, uid, *cols.values(), now))
        conn.execute("INSERT INTO application_events(app_id, at, from_status, to_status, source) "
                     "VALUES (?,?,NULL,'saved',?)", (cur.lastrowid, now, source))
    from . import jdsnap                                         # late: jdsnap imports this module
    jdsnap.on_save(conn, cur.lastrowid)                          # the JD text is kept from the first save
    return cur.lastrowid


def edit_field(conn, app_id: int, field: str, value: str | None, source: str = "app") -> None:
    """Edit one tracker column from the app (or, source='sheet', pulled from the Google Sheet); export will write
    the new value for that field."""
    if field not in FIELDS:
        raise ValueError(f"unknown tracker field {field}")
    row = conn.execute(f"SELECT edited, field_at, {field} AS old FROM applications WHERE id=?", (app_id,)).fetchone()
    if row is None:
        raise ValueError(f"no application {app_id}")
    edited = sorted(set(json.loads(row["edited"] or "[]")) | {field})
    now = db.iso(db.utcnow())
    stamps = row["field_at"] if source == "sheet" else _stamped(conn, app_id, field, now)
    with db.tx(conn):
        conn.execute(f"UPDATE applications SET {field}=?, edited=?, updated_at=?, field_at=? WHERE id=?",
                     (value, json.dumps(edited), now, stamps, app_id))
        history.record(conn, app_id, field, row["old"], value, "sheet" if source == "sheet" else "app")
        gstate.mark_dirty(conn)


def set_jd_text(conn, app_id: int, text: str, url: str | None = None) -> None:
    """The row's full JD (the hidden "JD Text" column) and, when the row has none, its posting link."""
    row = conn.execute("SELECT jd_text, job_url FROM applications WHERE id=?", (app_id,)).fetchone()
    now = db.iso(db.utcnow())
    with db.tx(conn):
        conn.execute("UPDATE applications SET jd_text=?, job_url=coalesce(nullif(job_url, ''), ?), updated_at=?, "
                     "field_at=? WHERE id=?", (text, url, now, _stamped(conn, app_id, "jd_text", now), app_id))
        history.record(conn, app_id, "jd_text", row["jd_text"], text, "app")
        if url and not row["job_url"]:
            history.record(conn, app_id, "job_url", row["job_url"], url, "app")
        gstate.mark_dirty(conn)


def followups_due(conn, profile: str) -> list[dict]:
    rows = conn.execute("SELECT * FROM followups_due WHERE profile=? ORDER BY applied_on DESC, id DESC", (profile,)).fetchall()
    return [dict(r) for r in rows]
