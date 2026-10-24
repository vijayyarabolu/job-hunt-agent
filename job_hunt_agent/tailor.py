"""The built-in resume tailor: what the Tailored resumes tab talks to when no resume-tailoring service URL is set.

It answers the same small HTTP-shaped API a separate service would (rgclient.request routes here in-process), so the
queue, chips, side panel and match scores work the same either way:

    GET  /api/jobs?profile=<p>[&lean=1]      -> {"jobs": [...]}
    POST /api/jobs {"jobs": [item, ...]}      -> {"ids": [...]}      item: profile, company, role, jd_text, jd_url,
                                                                     source_job_id, workplace, location
    GET  /api/jobs/<id>                        -> {"job": {...}, "jd_text": "..."}
    GET  /api/jobs/<id>/file/<docx|pdf>        -> the file
    POST /api/jobs/<id>/<action>               approve | accept | regenerate | fix | retry | cancel | archive | top
    POST /api/unarchive {"ids": [...]}         -> {"restored": [...]}

One background worker takes the queue in order. For each job Claude (the user's own Anthropic key) rewrites the
headline, summary, skills order and bullets of the base resume for that job description, with the rule that every
fact must already be in the resume: no new employers, titles, dates, numbers or tools. The answer keeps the same jobs
and the same number of bullets, so the panel can show a word-level diff against the base. Files are a .docx (read
back by tailored.py for the match score) and a plain PDF.

Records live in <home>/data/tailor/<id>.json; nothing leaves the machine except the Claude call.
"""
import difflib
import io
import json
import logging
import re
import threading
import time
from pathlib import Path

from . import atomic, claude_api, config, lexicon
from .resume import ats, baseline

log = logging.getLogger(__name__)
_LOCK = threading.RLock()
_WAKE = threading.Event()
_WORKER: list[threading.Thread] = []
ACTIONS = ("approve", "accept", "regenerate", "fix", "retry", "cancel", "archive", "top", "undo_edit", "gdoc",
           "reimport")
MAX_JD = 30_000
SYSTEM = ("You tailor a job seeker's resume to one job description. You may reword, reorder and emphasize, but every "
          "fact must already be in the resume you are given: never add an employer, title, date, number, metric, "
          "tool or skill that is not there. Plain, specific language; no buzzwords.")
PROMPT = """Rewrite this resume for the job below.

Rules:
- headline: the job's role title in the resume's own words (for example "Machine Learning Engineer").
- summary: 2 to 3 sentences, using the job's wording where the resume supports it.
- skills: the same categories and the same items, each category's items reordered so the ones the job asks for
  come first. Do not add or drop items.
- experience: the same jobs in the same order, each with exactly the same number of bullets. Rewrite a bullet only
  where the job's wording fits a fact the bullet already states; otherwise return it unchanged.
{note}
Resume (JSON):
{resume}

Job: {role} at {company}
Job description:
{jd}
"""
SCHEMA = {"type": "object", "required": ["headline", "summary", "skills", "experience"], "properties": {
    "headline": {"type": "string"}, "summary": {"type": "string"},
    "skills": {"type": "array", "items": {"type": "object", "required": ["category", "items"], "properties": {
        "category": {"type": "string"}, "items": {"type": "array", "items": {"type": "string"}}}}},
    "experience": {"type": "array", "items": {"type": "object", "required": ["company", "bullets"], "properties": {
        "company": {"type": "string"}, "bullets": {"type": "array", "items": {"type": "string"}}}}}}}


def folder() -> Path:
    return config.DATA / "tailor"


def _path(job_id: int) -> Path:
    return folder() / f"{int(job_id)}.json"


def _load(job_id: int) -> dict:
    try:
        return json.loads(_path(job_id).read_text())
    except (OSError, ValueError):
        raise KeyError(f"no tailored resume #{job_id}") from None


def _save(job: dict) -> None:
    job["updated_at"] = _now()
    atomic.write_text(_path(job["id"]), json.dumps(job, ensure_ascii=False))


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def _all() -> list[dict]:
    out = []
    for path in sorted(folder().glob("*.json"), key=lambda p: int(p.stem) if p.stem.isdigit() else 0):
        try:
            out.append(json.loads(path.read_text()))
        except (OSError, ValueError):
            continue
    return out


def _next_id() -> int:
    ids = [int(p.stem) for p in folder().glob("*.json") if p.stem.isdigit()]
    return max(ids, default=0) + 1


def base_resume(pid: str) -> dict:
    """The profile's base resume as a dict: baseline.json, else its first resume .docx parsed the same way."""
    home = config.PROFILES / pid
    try:
        return json.loads((home / "baseline.json").read_text())
    except (OSError, ValueError):
        pass
    for entry in config.profile(pid).get("resumes") or []:
        path = home / "resumes" / str(entry.get("file") or "")
        if path.is_file():
            return baseline.parse(path)
    raise ValueError("this profile has no resume yet: upload one in Settings > Profiles")


# ---- the API a resume service would answer ----

def handle(method: str, path: str, body: dict | None = None, raw: bool = False):
    """One request, answered in-process. KeyError = 404, ValueError = a refusal with its message."""
    path, _, query = path.partition("?")
    args = dict(p.split("=", 1) for p in query.split("&") if "=" in p)
    parts = [p for p in path.split("/") if p]
    with _LOCK:
        if parts == ["api", "jobs"] and method == "GET":
            lean = args.get("lean") == "1"
            return {"jobs": [_public(j, lean) for j in _all() if j.get("profile") == args.get("profile")]}
        if parts == ["api", "jobs"] and method == "POST":
            return {"ids": [_queue(item) for item in (body or {}).get("jobs") or []]}
        if parts == ["api", "unarchive"] and method == "POST":
            return {"restored": _unarchive((body or {}).get("ids") or [])}
        if parts == ["api", "google", "relink"]:
            return {"relinked": 0}
        if len(parts) == 3 and parts[:2] == ["api", "jobs"] and method == "GET":
            job = _load(int(parts[2]))
            return {"job": _public(job, False), "jd_text": job.get("jd_text") or ""}
        if len(parts) == 5 and parts[3] == "file" and method == "GET":
            return _file(_load(int(parts[2])), parts[4])
        if len(parts) == 4 and parts[:2] == ["api", "jobs"] and method == "POST":
            return _act(_load(int(parts[2])), parts[3], body or {})
    raise KeyError(path)


def _public(job: dict, lean: bool) -> dict:
    out = {k: v for k, v in job.items() if k not in ("jd_text", "tailored", "base_resume")}
    out["times"] = {"made": job.get("made_at"), "updated": job.get("updated_at")}
    if lean:
        out.pop("result", None)
    return out


def _queue(item: dict) -> int:
    pid = str(item.get("profile") or "")
    if pid not in config.profile_ids():
        raise ValueError(f"unknown profile {pid!r}")
    jd = str(item.get("jd_text") or "").strip()[:MAX_JD]
    if len(jd) < 80:
        raise ValueError("the job description is too short to tailor for")
    job = {"id": _next_id(), "profile": pid, "company": str(item.get("company") or "")[:200],
           "role": str(item.get("role") or "")[:200], "jd_text": jd, "jd_url": item.get("jd_url"),
           "source_job_id": item.get("source_job_id"), "location": item.get("location"),
           "workplace": item.get("workplace"), "status": "queued", "attempts": 0, "created_at": _now()}
    _save(job)
    _start_worker()
    _WAKE.set()
    return job["id"]


def _unarchive(ids) -> list[int]:
    back = []
    for job_id in ids:
        try:
            job = _load(int(job_id))
        except (KeyError, ValueError, TypeError):
            continue
        job["archived_at"] = None
        _save(job)
        back.append(job["id"])
    return back


def _act(job: dict, name: str, body: dict) -> dict:
    if name not in ACTIONS:
        raise KeyError(name)
    if name in ("gdoc", "reimport", "undo_edit"):
        raise ValueError("this needs a resume-tailoring service (Settings > Setup); the built-in tailor makes a "
                         ".docx you can edit in Word or Google Docs")
    if name in ("regenerate", "fix", "retry"):
        job.update(status="queued", error=None, note=(body.get("note") or job.get("note")))
        _save(job)
        _start_worker()
        _WAKE.set()
    elif name == "cancel":
        job.update(status="cancelled")
        _save(job)
    elif name == "archive":
        job["archived_at"] = _now()
        _save(job)
    elif name == "accept":
        job["status"] = "ready"
        _save(job)
    elif name == "approve":
        if job.get("status") != "ready":
            raise ValueError("only a ready resume can be approved")
        result = job.setdefault("result", {})
        result["approved"] = {"name": f"{_stem(job)}.docx", "at": _now()}
        _save(job)
    return {"ok": True, "status": job.get("status")}


# ---- the work ----

def _start_worker() -> None:
    if _WORKER and _WORKER[0].is_alive():
        return
    thread = threading.Thread(target=_work, daemon=True, name="resume-tailor")
    _WORKER[:] = [thread]
    thread.start()


def _work() -> None:
    while True:
        with _LOCK:
            todo = next((j for j in _all() if j.get("status") == "queued"), None)
            if todo:
                todo.update(status="generating", attempts=int(todo.get("attempts") or 0) + 1)
                _save(todo)
        if todo is None:
            _WAKE.wait(30)
            _WAKE.clear()
            if not any(j.get("status") == "queued" for j in _all()):
                return
            continue
        try:
            tailored = _tailor(todo)
            done = _finish(todo, tailored)
        except (claude_api.ClaudeError, ValueError, KeyError, OSError) as err:
            done = {**todo, "status": "failed", "error": str(err)[:300]}
        except Exception as err:                               # the worker must outlive one bad job
            log.exception("resume tailor: job %s failed", todo["id"])
            done = {**todo, "status": "failed", "error": f"unexpected error: {type(err).__name__}"}
        with _LOCK:
            current = _load(todo["id"])
            if current.get("status") == "generating":          # not cancelled meanwhile
                _save(done)


def _tailor(job: dict) -> dict:
    base = base_resume(job["profile"])
    if config.demo_mode():
        return demo_tailor(base, job)
    if not claude_api.configured():
        raise ValueError("built-in tailoring needs an Anthropic API key (Settings > Setup)")
    note = f"- The person's note: {job['note']}\n" if job.get("note") else ""
    prompt = PROMPT.format(note=note, resume=json.dumps({k: base.get(k) for k in ("headline", "summary", "skills",
                                                                                  "experience")}, ensure_ascii=False),
                           role=job["role"], company=job["company"], jd=job["jd_text"])
    out, usage = claude_api.ask(prompt, SYSTEM, SCHEMA, effort="medium")
    job["usage"] = usage
    return _keep_shape(base, out)


def demo_tailor(base: dict, job: dict) -> dict:
    """Demo mode: no Claude. Skills the job names move to the front and the headline follows the role."""
    wanted = {t.name for t in ats.extract(job["jd_text"])}
    skills = [{"category": s["category"], "items": sorted(s["items"], key=lambda i: not (set(lexicon.find(i)) & wanted))}
              for s in base.get("skills") or []]
    summary = base.get("summary") or ""
    return {**base, "headline": job["role"] or base.get("headline", ""), "skills": skills,
            "summary": f"{job['role']} candidate. {summary}".strip() if job.get("role") else summary}


def _keep_shape(base: dict, out: dict) -> dict:
    """Claude's answer forced back into the base resume's shape: the same jobs and bullet counts, the same skill
    items (reordered only); anything missing keeps the base text."""
    skills = []
    answered = {s.get("category"): s.get("items") or [] for s in out.get("skills") or [] if isinstance(s, dict)}
    for cat in base.get("skills") or []:
        items = list(cat["items"])
        order = [i for i in answered.get(cat["category"], []) if i in items]
        skills.append({"category": cat["category"], "items": order + [i for i in items if i not in order]})
    jobs = []
    new_jobs = [j for j in out.get("experience") or [] if isinstance(j, dict)]
    for n, job in enumerate(base.get("experience") or []):
        got = new_jobs[n].get("bullets") if n < len(new_jobs) else None
        bullets = [str(got[i]).strip() if got and i < len(got) and str(got[i]).strip() else b
                   for i, b in enumerate(job.get("bullets") or [])]
        jobs.append({**job, "bullets": bullets})
    return {**base, "headline": str(out.get("headline") or base.get("headline") or "").strip(),
            "summary": str(out.get("summary") or base.get("summary") or "").strip(), "skills": skills,
            "experience": jobs}


def _ops(a: str, b: str) -> dict:
    """A word-level diff in the panel's format: {"changed", "ops": [[op, base, new], ...], "base", "new"}."""
    left, right = a.split(), b.split()
    ops = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(a=left, b=right, autojunk=False).get_opcodes():
        ops.append([tag, " ".join(left[i1:i2]) or None, " ".join(right[j1:j2]) or None])
    return {"changed": a.strip() != b.strip(), "ops": ops, "base": a, "new": b}


def diff(base: dict, new: dict) -> dict:
    jobs, changed, total = [], 0, 0
    for old, now in zip(base.get("experience") or [], new.get("experience") or []):
        bullets = [_ops(a, b) for a, b in zip(old.get("bullets") or [], now.get("bullets") or [])]
        changed += sum(b["changed"] for b in bullets)
        total += len(bullets)
        jobs.append({"company": old.get("company"), "title": old.get("title"), "dates": old.get("dates"),
                     "bullets": bullets})
    skills = [{"category": a["category"], "items": _ops(", ".join(a["items"]), ", ".join(b["items"]))}
              for a, b in zip(base.get("skills") or [], new.get("skills") or [])]
    return {"headline": _ops(base.get("headline") or "", new.get("headline") or ""),
            "summary": _ops(base.get("summary") or "", new.get("summary") or ""), "skills": skills, "jobs": jobs,
            "changed_bullets": changed, "total_bullets": total}


def _ats(jd: str, base: dict, new: dict) -> dict:
    terms = ats.extract(jd)
    before = ats.coverage(terms, set(ats.resume_terms(base)))
    after = ats.coverage(terms, set(ats.resume_terms(new)))
    return {"base_pct": round(100 * before.pct), "tailored_pct": round(100 * after.pct),
            "label": "Keywords from the job found in the resume text", "missing": list(after.missing)[:12]}


def _finish(job: dict, tailored: dict) -> dict:
    base = base_resume(job["profile"])
    folder().mkdir(parents=True, exist_ok=True)
    (folder() / f"{job['id']}.docx").write_bytes(to_docx(job["profile"], tailored))
    return {**job, "status": "ready", "error": None, "made_at": _now(), "tailored": tailored,
            "result": {"base": {"id": "main", "label": base.get("headline") or "Base resume"},
                       "headline": tailored.get("headline"), "diff": diff(base, tailored),
                       "ats": _ats(job["jd_text"], base, tailored), "files": {"pdf": True, "docx": True}}}


def _stem(job: dict) -> str:
    name = str(config.profile(job["profile"]).get("name") or "Resume").replace(" ", "_")
    role = re.sub(r"[^A-Za-z0-9]+", "_", job.get("role") or "").strip("_")[:60]
    return f"{name}_{role}_Resume" if role else f"{name}_Resume"


def _file(job: dict, which: str) -> tuple[bytes, str]:
    if job.get("status") != "ready" or not job.get("tailored"):
        raise KeyError("this resume is not ready")
    if which == "docx":
        return (folder() / f"{job['id']}.docx").read_bytes(), \
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if which == "pdf":
        return to_pdf(job["profile"], job["tailored"]), "application/pdf"
    raise KeyError(which)                                      # no page images: the panel shows the PDF


# ---- files ----

def _contact(pid: str) -> tuple[str, str]:
    profile = config.profile(pid)
    bits = [profile.get(k) for k in ("location", "email", "phone", "linkedin", "github")]
    return str(profile.get("name") or ""), " | ".join(str(b) for b in bits if b)


def to_docx(pid: str, resume: dict) -> bytes:
    """A plain .docx in the layout baseline.parse reads back (headline, Summary, Skills, Work Experience)."""
    import docx
    doc = docx.Document()
    name, contact = _contact(pid)
    doc.add_heading(name, level=0)
    if resume.get("headline"):
        doc.add_paragraph(resume["headline"])
    if contact:
        doc.add_paragraph(contact)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(resume.get("summary") or "")
    doc.add_heading("Skills", level=1)
    for cat in resume.get("skills") or []:
        doc.add_paragraph(f"{cat['category']}: {', '.join(cat['items'])}")
    doc.add_heading("Work Experience", level=1)
    for job in resume.get("experience") or []:
        doc.add_paragraph(f"{job.get('company') or ''}\t{job.get('dates') or ''}")
        if job.get("title"):
            doc.add_paragraph(job["title"])
        for bullet in job.get("bullets") or []:
            doc.add_paragraph(bullet, style="List Bullet")
    if resume.get("education"):
        doc.add_heading("Education", level=1)
        for line in resume["education"]:
            doc.add_paragraph(line if isinstance(line, str) else " ".join(str(v) for v in line.values() if v))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _lines(pid: str, resume: dict) -> list[tuple[str, str]]:
    """(style, text) lines for the PDF: h = heading, b = bullet, t = text, n = name."""
    name, contact = _contact(pid)
    out = [("n", name), ("t", resume.get("headline") or ""), ("t", contact), ("h", "Summary"),
           ("t", resume.get("summary") or ""), ("h", "Skills")]
    out += [("t", f"{c['category']}: {', '.join(c['items'])}") for c in resume.get("skills") or []]
    out.append(("h", "Work Experience"))
    for job in resume.get("experience") or []:
        out.append(("t", f"{job.get('company') or ''}    {job.get('dates') or ''}".strip()))
        if job.get("title"):
            out.append(("t", job["title"]))
        out += [("b", b) for b in job.get("bullets") or []]
    return [(s, t) for s, t in out if t]


def _wrap(text: str, width: int) -> list[str]:
    words, lines, line = text.split(), [], ""
    for word in words:
        if len(line) + len(word) + 1 > width and line:
            lines.append(line)
            line = word
        else:
            line = f"{line} {word}".strip()
    return lines + ([line] if line else [])


def _pdf_text(text: str) -> str:
    plain = text.encode("latin-1", "replace").decode("latin-1")
    return plain.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def to_pdf(pid: str, resume: dict) -> bytes:
    """A plain one-column PDF (Helvetica, US Letter), written by hand so no PDF library is needed."""
    sizes = {"n": 16, "h": 12, "t": 10, "b": 10}
    pages, ops, y = [], [], 750
    for style, text in _lines(pid, resume):
        size = sizes[style]
        if style == "h":
            y -= 6
        for i, line in enumerate(_wrap(text, 92 if size == 10 else 70)):
            if y < 60:
                pages.append(ops)
                ops, y = [], 750
            x = 72 if style != "b" else (72 if i == 0 else 82)
            label = f"- {line}" if style == "b" and i == 0 else line
            font = "F2" if style in ("n", "h") else "F1"
            ops.append(f"BT /{font} {size} Tf {x} {y} Td ({_pdf_text(label)}) Tj ET")
            y -= size + 4
    pages.append(ops)
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None,
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold >>"]
    kids = []
    for page_ops in pages:
        stream = "\n".join(page_ops).encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode("latin-1") + stream + b"\nendstream")
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {len(objects)} 0 R "
                       "/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> >>")
        kids.append(len(objects))
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] /Count {len(kids)} >>"
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for n, obj in enumerate(objects, 1):
        offsets.append(out.tell())
        body = obj if isinstance(obj, bytes) else obj.encode("latin-1")
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    out.write("".join(f"{o:010d} 00000 n \n" for o in offsets).encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()
