"""Each person's tracker columns, exactly as their real sheet has them, plus the app's own fields kept apart.

One Layout per profile, read from the person's real workbook (its header rows are the truth for names and order):
  - plain profile (Alex): `tracker_template` in profile.yaml (default data/import/Application_Tracker.xlsx), first
    tab, plus `tracker_add_columns: [{header: Address, after: Role Title}]` for columns the app adds to the real sheet;
  - mapped profile (Jordan): profiles/<p>/tracker_map.yaml (her tabs in her column order, her vendors tab).
A column with no app field (Jordan's "Column 1") is a pass-through: kept from the imported cell, read-only.
`cells()` is the one place a row's sheet values come from, so the tracker view, the Excel export and the Google Sheet
sync always show the same thing. APP_EXTRAS are the app's own fields, shown and exported in a separate section.
"""
import json
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import openpyxl

from . import config, tracker
from .gsync_values import JD_TEXT_HEADER, MAX_CELL
from .tracker import COLUMNS, FIELDS

MAPPING_FILE = "tracker_map.yaml"
ADDED_FIELD = {"Address": "address", "Addresss": "address", "Team Name": "team_name"}
EXTRAS_SHEET = "job-hunt-agent app fields"          # the separate, labeled section in Excel exports
APP_EXTRAS = (("job_url", "Job URL"), ("apply_url", "Apply URL"), ("salary", "Salary"), ("workplace", "Workplace"),
              ("status", "Pipeline status"), ("resume_used", "Resume used"), ("match_pct", "Match %"),
              ("fu1_due", "Follow-up 1 due"), ("fu2_due", "Follow-up 2 due"), ("notes", "Notes"),
              ("location", "Location"), ("employment_type", "Employment type"), ("source", "Source"),
              ("updated_at", "Last updated"))


@dataclass(frozen=True)
class Column:
    header: str
    field: str | None            # applications column; None = pass-through cell (read-only)
    target: str | None = None    # mapped profiles: the 14-column tracker name this header stands for


# GOAL6 JD format: a hidden column the app adds at the far right of each Google Sheet tab; synced like a column but
# never part of a layout, an import or an Excel export (the JD cell holds the linked title, this one the full text).
JD_TEXT = Column(JD_TEXT_HEADER, "jd_text")


@dataclass(frozen=True)
class Tab:
    name: str
    kind: str                    # apps | vendors
    columns: tuple[Column, ...]
    index: int = 1               # apps tabs: 1-based position among the person's application tabs

    @property
    def headers(self) -> list[str]:
        return [c.header for c in self.columns]


@dataclass(frozen=True)
class Layout:
    profile: str
    tabs: tuple[Tab, ...]
    template: Path | None
    mapping: Path | None         # tracker_map.yaml for a mapped profile
    added: tuple[tuple[str, str], ...] = ()   # (header, after) columns the app adds to the template

    @property
    def app_tabs(self) -> tuple[Tab, ...]:
        return tuple(t for t in self.tabs if t.kind == "apps")


@lru_cache(maxsize=16)
def _headers(path: str, mtime: float) -> dict[str, list[str]]:
    book = openpyxl.load_workbook(path, read_only=True)
    try:
        out = {}
        for sheet in book.worksheets:
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            header = [str(v).strip() if v is not None else "" for v in first]
            while header and not header[-1]:
                header.pop()
            out[sheet.title] = header
        return out
    finally:
        book.close()


def template_headers(path: Path) -> dict[str, list[str]]:
    return _headers(str(path), path.stat().st_mtime)


def _plain(profile: dict, template: Path | None = None) -> Layout:
    template = template or (Path(config.ROOT / profile["tracker_template"]) if profile.get("tracker_template")
                            else tracker.default_import_path())
    if template.exists():
        name, header = next(iter(template_headers(template).items()))
    else:
        name, header = "Job_Application_Tracker", list(COLUMNS)
    known = dict(zip(COLUMNS, FIELDS)) | ADDED_FIELD
    columns = [Column(h, known.get(h)) for h in header]
    added = tuple((str(a["header"]), str(a["after"])) for a in profile.get("tracker_add_columns") or ())
    for head, after in added:
        if head not in header:
            at = next((i + 1 for i, c in enumerate(columns) if c.header == after), len(columns))
            columns.insert(at, Column(head, known.get(head)))
    return Layout(profile["id"], (Tab(name, "apps", tuple(columns)),), template if template.exists() else None, None,
                  added)


def of_mapping(pid: str, mapping_path: Path) -> Layout:
    """A mapped profile's layout straight from its tracker_map.yaml (no profile.yaml needed)."""
    return _mapped({"id": pid}, mapping_path)


def _mapped(profile: dict, mapping_path: Path) -> Layout:
    from . import tracker_map
    mapping = tracker_map.load_mapping(mapping_path)
    template = config.ROOT / mapping.template if mapping.template else None
    headers = template_headers(template) if template and template.exists() else {}
    tabs = []
    for index, spec in enumerate(mapping.sheets, 1):
        targets = {**spec.get("columns", {})}
        order = headers.get(spec["name"]) or [*targets, *spec.get("fill_if_empty", {})]
        cols = []
        for head in order:
            target = targets.get(head)
            field = FIELDS[COLUMNS.index(target)] if target in COLUMNS else ADDED_FIELD.get(target or "")
            cols.append(Column(head, field, target))
        tabs.append(Tab(spec["name"], "apps", tuple(cols), index))
    if mapping.vendors:
        spec = mapping.vendors
        order = headers.get(spec["sheet"]) or list(spec["columns"])
        tabs.append(Tab(spec["sheet"], "vendors", tuple(Column(h, spec["columns"].get(h)) for h in order)))
    return Layout(profile["id"], tuple(tabs), template if template and template.exists() else None, mapping_path)


def of_profile(pid: str, template: Path | None = None) -> Layout:
    """The profile's layout; `template` (plain profiles) reads the base columns from that workbook instead."""
    profile = config.profile(pid)
    mapping = config.PROFILES / pid / MAPPING_FILE
    return _mapped(profile, mapping) if mapping.exists() else _plain(profile, Path(template) if template else None)


def tab_of(layout: Layout, app: dict) -> Tab:
    """The application tab a row belongs to: its imported sheet (mapped import row numbers), else the first."""
    from . import tracker_map
    tabs = layout.app_tabs
    return tabs[tracker_map._sheet_index(app, len(tabs)) - 1] if layout.mapping else tabs[0]


def _raw(app: dict) -> list:
    return json.loads(app["raw_json"]) if app.get("raw_json") else []


def _passthrough(app: dict, header: str):
    tagged = next((c for c in _raw(app) if isinstance(c, dict) and c.get("col") == header), None)
    return (tracker.decode(tagged), tagged.get("h")) if tagged else (None, None)


def _as_text(value) -> bool:
    return isinstance(value, str) and value.startswith(tracker._FORMULA_CHARS)


def cell(app: dict, column: Column, layout: Layout, links: dict | None = None) -> tuple:
    """(value, link, write_as_text) of one sheet cell of one application, exactly as the Excel export writes it
    (an imported formula stays a formula; typed text that looks like one stays text)."""
    from . import tracker_map
    if column.field in FIELDS:
        idx = FIELDS.index(column.field)
        value, as_text = tracker._cell_value(app, idx)
        if column.field == "status_ar":
            value = tracker._status_cell(value, app)
            as_text = _as_text(value)
        return value, tracker._cell_link(app, idx, links or {}), as_text
    if column == JD_TEXT:                                # cut like the Sheet cuts it, so both sides compare equal
        return (app.get("jd_text") or None) and app["jd_text"][:MAX_CELL], None, True
    if column.field:                                    # Address / Team Name
        value, link = tracker_map._extra_cell(app, column.target or column.header) if layout.mapping \
            else (app.get(column.field), None)
    else:
        value, link = _passthrough(app, column.header)
    return value, link, _as_text(value)


def cells(app: dict, layout: Layout) -> tuple[Tab, list]:
    tab = tab_of(layout, app)
    return tab, [cell(app, c, layout)[0] for c in tab.columns]


def write_extras(book, rows: list[tuple[str, int, dict]]) -> None:
    """The separate, labeled 'job-hunt-agent app fields' sheet: one line per exported row, keyed by (tab, sheet row) so an
    import puts the values back on the row it just read (also into an empty database)."""
    if EXTRAS_SHEET in book.sheetnames:
        del book[EXTRAS_SHEET]
    sheet = book.create_sheet(EXTRAS_SHEET)
    sheet.append(["Sheet tab", "Sheet row", "Row ID", "Company Name", "Role Title", *(label for _, label in APP_EXTRAS)])
    for tab, row, values in rows:
        line = [tab, row, values.get("id"), values.get("company"), values.get("role_title"),
                *(values.get(k) for k, _ in APP_EXTRAS)]
        sheet.append(line)
        for out in sheet[sheet.max_row]:
            if _as_text(out.value):
                out.data_type, out.quotePrefix = "s", True


IMPORTED_EXTRAS = ("job_url", "apply_url", "salary", "workplace", "resume_used", "match_pct", "notes", "location",
                   "employment_type", "source")


def read_extras(book) -> dict[tuple[str, int], dict]:
    """(tab, sheet row) -> app-only field values from a 'job-hunt-agent app fields' sheet (empty when there is none)."""
    if EXTRAS_SHEET not in book.sheetnames:
        return {}
    sheet = book[EXTRAS_SHEET]
    header = [str(c.value or "").strip() for c in sheet[1]]
    label = {lbl: key for key, lbl in APP_EXTRAS}
    out = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        if rec.get("Sheet tab") is None or rec.get("Sheet row") is None:
            continue
        out[(str(rec["Sheet tab"]), int(rec["Sheet row"]))] = {
            label[h]: v for h, v in rec.items() if label.get(h) in IMPORTED_EXTRAS and v not in (None, "")}
    return out


def apply_extras(conn, profile: str, found: dict[tuple[str, int], dict], row_number) -> int:
    """Put app-only fields back from an export's extras sheet onto the rows the import matched (`row_number(tab, row)`
    gives the stored xlsx_row). Never clears a value; returns how many rows changed."""
    changed = 0
    for (tab, row), values in found.items():
        app = conn.execute("SELECT * FROM applications WHERE profile = ? AND xlsx_row = ?",
                           (profile, row_number(tab, row))).fetchone()
        if app is None:
            continue
        sets = {k: v for k, v in values.items() if str(app[k] if app[k] is not None else "") != str(v)}
        if sets:
            conn.execute(f"UPDATE applications SET {', '.join(f'{k} = ?' for k in sets)} WHERE id = ?",
                         [*sets.values(), app["id"]])
            changed += 1
    return changed


def extras(app: dict, dues: dict | None = None) -> dict:
    """The app's own fields (never mixed into the sheet columns)."""
    dues = dues or {}
    out = {k: app.get(k) for k, _ in APP_EXTRAS if k not in ("fu1_due", "fu2_due")}
    out["fu1_due"] = (dues.get("fu1") or {}).get("due")
    out["fu2_due"] = (dues.get("fu2") or {}).get("due")
    return out


def describe(layout: Layout) -> dict:
    """The layout for the UI: every tab's headers in order, and the app-fields section."""
    return {"tabs": [{"name": t.name, "kind": t.kind, "index": t.index,
                      "columns": [{"header": c.header, "field": c.field} for c in t.columns]} for t in layout.tabs],
            "extras": [{"field": k, "label": v} for k, v in APP_EXTRAS], "extras_label": EXTRAS_SHEET}


# ---------- adding a column to a copy of the real workbook (Excel export) ----------

def _shift_range(ref: str, at: int) -> str:
    """A1-style range after inserting a column at 1-based `at`: ranges right of it move, ranges across it widen."""
    from openpyxl.worksheet.cell_range import CellRange
    rng = CellRange(ref)
    if rng.min_col >= at:
        rng.shift(col_shift=1)
    elif rng.max_col >= at:
        rng.expand(right=1)
    return rng.coord


def _shift_widths(sheet, at: int, width: float | None) -> None:
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.dimensions import ColumnDimension, DimensionHolder
    spans = [(d.min, d.max, d.width, d.hidden, d.customWidth) for d in sheet.column_dimensions.values() if d.min]
    holder = DimensionHolder(worksheet=sheet)
    for lo, hi, w, hidden, custom in spans:
        lo2, hi2 = (lo + 1, hi + 1) if lo >= at else (lo, hi + 1 if hi >= at else hi)
        holder[get_column_letter(lo2)] = ColumnDimension(sheet, min=lo2, max=hi2, width=w, hidden=hidden,
                                                         customWidth=custom)
    if width and not any(lo2 <= at <= hi2 for lo2, hi2 in ((d.min, d.max) for d in holder.values())):
        holder[get_column_letter(at)] = ColumnDimension(sheet, min=at, max=at, width=width, customWidth=True)
    sheet.column_dimensions = holder


def _shift_conditional(sheet, at: int) -> None:
    import re
    from openpyxl.formatting.formatting import ConditionalFormattingList
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.cell_range import CellRange
    old, sheet.conditional_formatting = sheet.conditional_formatting, ConditionalFormattingList()
    for cf in old:
        refs = [r.coord for r in cf.sqref.ranges]
        first = CellRange(refs[0])
        new_refs = [_shift_range(r, at) for r in refs]
        new_first = CellRange(new_refs[0])
        before = f"{get_column_letter(first.min_col)}{first.min_row}"
        after = f"{get_column_letter(new_first.min_col)}{new_first.min_row}"
        for rule in cf.rules:
            if before != after:      # relative formulas are anchored on the range's top-left cell
                rule.formula = [re.sub(rf"(?<![A-Z$]){before}(?!\d)", after, f) for f in rule.formula or []]
            sheet.conditional_formatting.add(" ".join(new_refs), rule)


def insert_column(sheet, at: int, header: str, width: float | None = None) -> None:
    """Insert `header` as column `at` (1-based) in a copy of the real workbook, keeping what Sheets keeps when a column
    is inserted: data validation and conditional formatting ranges (and the CF formula anchor), column widths, the
    header cell's style, and the Excel table (its range and its column list)."""
    from copy import copy
    from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
    from openpyxl.worksheet.table import TableColumn
    sheet.insert_cols(at)
    for dv in sheet.data_validations.dataValidation:
        dv.sqref = MultiCellRange(" ".join(_shift_range(r.coord, at) for r in dv.sqref.ranges))
    _shift_conditional(sheet, at)
    _shift_widths(sheet, at, width)
    for table in sheet.tables.values():
        span = CellRange(table.ref)
        if not span.min_col < at <= span.max_col + 1:
            continue                                     # the table does not touch the new column
        table.ref = _shift_range(table.ref, at) if at <= span.max_col else CellRange(
            min_col=span.min_col, min_row=span.min_row, max_col=span.max_col + 1, max_row=span.max_row).coord
        cols = list(table.tableColumns)
        cols.insert(at - span.min_col, TableColumn(id=max(c.id for c in cols) + 1, name=header))
        table.tableColumns = cols
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    left = sheet.cell(row=1, column=at - 1) if at > 1 else None
    head = sheet.cell(row=1, column=at, value=header)
    if left is not None and left.has_style:
        head._style = copy(left._style)
