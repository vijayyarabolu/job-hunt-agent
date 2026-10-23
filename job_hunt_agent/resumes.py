"""A profile's resumes (decision 007): `resumes: [{id, label, file, families}]` in profile.yaml, files in
profiles/<pid>/resumes/. Each resume's searchable text is the whole .docx (paragraphs and tables, like a plain-text
keyword tool), parsed once per file content (md5). No `resumes:` list -> the single baseline.json is resume "main",
so a profile keeps working until its resumes are dropped in; adding them later needs no code change and no command.
"""
import hashlib
import io
import json
import logging
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import docx

from . import config, lexicon, resume_facts
from .resume import ats

log = logging.getLogger(__name__)
MAIN_ID = "main"
RESUME_DIR = "resumes"
DEFAULT_FILE_PATTERN = "{name}_{title}_Resume({company}).pdf"
_UNSAFE_FILE_CHARS = re.compile(r'[\\/:*?"<>|\n\r\t]+')


@dataclass(frozen=True)
class Resume:
    id: str
    label: str
    file: str | None              # file name under profiles/<pid>/resumes/; None for the baseline.json fallback
    families: tuple[str, ...]
    terms: frozenset[str]         # canonical lexicon names found in the resume text
    digest: str                   # md5 of the source bytes (cache key and hash input)
    text: str = ""                # lowercased, whitespace-collapsed text (decision 010: literal matching)
    years: float | None = None    # years covered by its experience date ranges (resume_facts.years_of)
    risks: tuple[str, ...] = ()   # layout parsing risks found in the file (resume_facts.parse_risks)
    size_mb: float | None = None  # file size; None for the baseline.json fallback (no file to check)


@lru_cache(maxsize=32)
def _docx_text(digest: str, data: bytes) -> str:
    document = docx.Document(io.BytesIO(data))
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


@lru_cache(maxsize=32)
def _terms(digest: str, text: str) -> frozenset[str]:
    return frozenset(lexicon.find(text))


def _from_file(pid: str, entry: dict) -> Resume | None:
    path = config.PROFILES / pid / RESUME_DIR / str(entry["file"])
    try:
        data = path.read_bytes()
    except OSError as err:
        log.warning("resume %s/%s unreadable: %s", pid, entry.get("id"), err)
        return None
    digest = hashlib.md5(data).hexdigest()
    try:
        text = _docx_text(digest, data)
    except Exception as err:          # a corrupt .docx must not take the profile's scores down
        log.warning("resume %s is not a readable .docx: %s", path, err)
        return None
    return Resume(str(entry["id"]), str(entry.get("label") or entry["id"]), str(entry["file"]),
                  tuple(entry.get("families") or ()), _terms(digest, text), digest, normalized(text),
                  resume_facts.years_of(text), _risks(digest, data, str(entry["file"])), resume_facts.size_mb(data))


def _from_baseline(profile: dict) -> Resume | None:
    path: Path = config.PROFILES / profile["id"] / "baseline.json"
    if not path.exists():
        return None
    raw = path.read_bytes()
    digest = hashlib.md5(raw).hexdigest()
    resume = json.loads(raw)
    text = ats.resume_text(resume)
    dates = "\n".join(str(job.get("dates") or "") for job in resume.get("experience", []))
    return Resume(MAIN_ID, str(profile.get("headline_default") or "Resume"), None, tuple(profile.get("families") or ()),
                  _terms(digest, text), digest, normalized(text), resume_facts.years_of(dates))


def normalized(text: str) -> str:
    """The text literal matching searches: lowercase, every run of whitespace one space."""
    return " ".join((text or "").lower().split())


@lru_cache(maxsize=32)
def _risks(digest: str, data: bytes, file: str) -> tuple[str, ...]:
    return resume_facts.parse_risks(data, file)


def validate(entries) -> list[dict]:
    """The `resumes:` list, checked: a list of mappings with unique id and a file; bad entries raise ValueError."""
    if not isinstance(entries, list):
        raise ValueError("resumes: must be a list of {id, label, file, families}")
    seen = set()
    for e in entries:
        if not isinstance(e, dict) or not e.get("id") or not e.get("file"):
            raise ValueError(f"resumes: entry needs id and file: {e!r}")
        if e["id"] in seen:
            raise ValueError(f"resumes: duplicate id {e['id']!r}")
        if Path(str(e["file"])).name != str(e["file"]):
            raise ValueError(f"resumes: file must be a plain file name in {RESUME_DIR}/: {e['file']!r}")
        seen.add(e["id"])
    return entries


def of_profile(profile: dict) -> tuple[Resume, ...]:
    """The profile's resumes in profile.yaml order (unreadable files are skipped with a warning)."""
    entries = profile.get("resumes")
    if not entries:
        main = _from_baseline(profile)
        return (main,) if main else ()
    loaded = (_from_file(profile["id"], e) for e in validate(entries))
    return tuple(r for r in loaded if r is not None)


def for_families(resumes: tuple[Resume, ...], families: list[str]) -> Resume | None:
    """The title-mapped resume: the first whose families include one of the job's families (job order first)."""
    for fam in families:
        for r in resumes:
            if fam in r.families:
                return r
    return None


def base_file_name(profile: dict, resume_id: str | None = None, ext: str = ".pdf",
                   role: str | None = None) -> str | None:
    """The base resume's own file name, e.g. Alex_Rivera_AI_Engineer_Resume.pdf / Jordan_Lee_AI_Engineer.pdf:
    every resume file (download, Drive, Sheet cell) is named like the person's real resumes (user, 2026-10-07).
    The profile.yaml `resumes:` entry with that id, else the only one listed, else the only PDF in resumes/.
    With `role`, the title in the name follows the job's role, not the base it was built from (for_role)."""
    entries = [e for e in (profile.get("resumes") or []) if isinstance(e, dict) and e.get("file")]
    entry = next((e for e in entries if resume_id and str(e.get("id")) == str(resume_id)), None)
    entry = entry or (entries[0] if len(entries) == 1 else None)
    stem = None
    if entry:
        stem = Path(str(entry["file"])).stem
    elif not entries and profile.get("id"):
        folder = config.PROFILES / str(profile["id"]) / "resumes"
        pdfs = sorted(folder.glob("*.pdf")) if folder.is_dir() else []
        stem = pdfs[0].stem if len(pdfs) == 1 else None
    return for_role(profile, stem, role) + ext if stem else None


# The name's title follows the job's role, never the template it was tailored from (user 2026-10-15: a Waymo
# "Senior Machine Learning Engineer" tailored from the Data Scientist base is Alex_Rivera_ML_Engineer_Resume).
# Same families, same order as resumegen's headline.for_job: data science, data engineering, AI+ML, ML, AI.
def _words(*patterns: str) -> re.Pattern:
    return re.compile(r"(?<![a-z0-9])(" + "|".join(patterns) + r")(?![a-z0-9])", re.I)


_DS = _words(r"data scien\w*", r"decision scien\w*", r"analytics scien\w*", r"statistician")
_DE = _words(r"data engineer\w*", r"analytics engineer\w*", r"data platforms?", r"etl", r"elt", r"data pipelines?",
             r"data warehous\w*", r"data infrastructure", r"big data")
_ML = _words(r"ml", r"machine[\s-]learning", r"mlops", r"applied scientists?", r"deep[\s-]learning",
             r"computer vision", r"nlp", r"natural language processing")
_AI = _words(r"ai", r"a\.i\.?", r"artificial intelligence", r"gen\s?ai", r"generative", r"llms?", r"llmops",
             r"agentic", r"ai tools", r"applied ai")


def role_title(role: str | None, allowed=None) -> str | None:
    """The file-name title for a posting (or a resumegen headline): ML_Engineer, AI_Engineer, AI_ML_Engineer,
    Data_Scientist, Data_Engineer; None when the role says none of them (the base's own title stays). `allowed`
    (profile.yaml resume_titles, e.g. three) maps a title outside it to its first one, as resumegen's
    headline does."""
    token = _role_token(role)
    return allowed[0] if token and allowed and token not in allowed else token


def _role_token(role: str | None) -> str | None:
    text = (role or "").replace("\u00a0", " ")
    if not text.strip():
        return None
    if _DS.search(text):
        return "Data_Scientist"
    if _DE.search(text):
        return "Data_Engineer"
    has_ai, has_ml = bool(_AI.search(text)), bool(_ML.search(text))
    if has_ai and has_ml:
        return "AI_ML_Engineer"
    if has_ml:
        return "ML_Engineer"
    if has_ai:
        return "AI_Engineer"
    return None


def title_label(role: str | None) -> str | None:
    """role_title for people: "ML Engineer", "AI/ML Engineer", "Data Scientist"."""
    t = role_title(role)
    return t.replace("AI_ML", "AI/ML").replace("_", " ") if t else None


def for_role(profile: dict, stem: str, role: str | None) -> str:
    """`stem` (a base resume's name) with its title swapped for the role's: the person's own base with that title
    when they have one (Alex_Rivera_ML_Engineer_Resume), else their naming style with the new title
    (Jordan_Lee_AI_Engineer_Resume -> Jordan_Lee_ML_Engineer_Resume). Unknown role or person: `stem` unchanged."""
    token = role_title(role, profile.get("resume_titles"))
    ends_with = re.compile(rf"(?:^|_)(?<!AI_){token}(?:_Resume)?$", re.I)   # ML_Engineer never matches AI_ML_Engineer
    if not token or ends_with.search(stem):
        return stem
    for e in profile.get("resumes") or []:
        other = Path(str(e.get("file") or "")).stem if isinstance(e, dict) else ""
        if other and ends_with.search(other):
            return other
    person = re.sub(r"\s+", "_", str(profile.get("name") or "").strip())
    if not person or not stem.lower().startswith(person.lower() + "_"):
        return stem
    suffix = "_Resume" if stem.lower().endswith("_resume") else ""
    return f"{stem[:len(person)]}_{token}{suffix}"


def file_name(profile: dict, resume: Resume | None, company: str | None, role: str | None = None) -> str:
    """Tracker Resume cell: the base resume's own name + the company, e.g. Alex_Rivera_ML_Engineer_Resume(Acme).pdf
    (user, 2026-10-08: easy to find later); the profile.yaml resume_file_pattern only when no base file is known."""
    firm = _UNSAFE_FILE_CHARS.sub(" ", (company or "").strip()).strip()
    tail = f"({firm}).pdf" if firm else ".pdf"
    if resume is not None and resume.file:
        return for_role(profile, Path(resume.file).stem, role) + tail
    known = base_file_name(profile, ext="", role=role)
    if known:
        return known + tail
    pattern = str(profile.get("resume_file_pattern") or DEFAULT_FILE_PATTERN)
    title = (resume.label if resume else str(profile.get("headline_default") or "Resume")).strip().replace(" ", "_")
    name = str(profile.get("name") or profile["id"]).strip().replace(" ", "_")
    company = _UNSAFE_FILE_CHARS.sub(" ", (company or "").strip()).strip()
    return pattern.format(name=name, title=title, company=company)
